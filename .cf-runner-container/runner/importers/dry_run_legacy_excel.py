from __future__ import annotations

import argparse
import json
import re
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def now_iso() -> str:
    return datetime.now().isoformat(timespec="seconds")


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat()
    return str(value).strip()


def parse_number(value: Any) -> float | None:
    text = clean_text(value)
    if not text:
        return None
    match = re.search(r"-?\d+(?:\.\d+)?", text)
    if not match:
        return None
    try:
        return float(match.group(0))
    except ValueError:
        return None


def normalize_shift(value: Any) -> str | None:
    text = clean_text(value)
    if not text:
        return None
    if "白" in text:
        return "day"
    if "夜" in text:
        return "night"
    return None


def split_deviation(value: Any) -> tuple[str | None, float | None]:
    text = clean_text(value)
    if not text:
        return (None, None)
    direction = None
    if "左" in text:
        direction = "left"
    elif "右" in text:
        direction = "right"
    elif "上" in text:
        direction = "up"
    elif "下" in text:
        direction = "down"
    return (direction, parse_number(text))


def find_value(row: dict[str, Any], aliases: list[str]) -> Any:
    for key, value in row.items():
        key_text = clean_text(key)
        if not key_text:
            continue
        for alias in aliases:
            if alias in key_text:
                if clean_text(value):
                    return value
    return None


def read_rows(file_path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(file_path, data_only=True)
    sheet = workbook.active
    raw_headers = [clean_text(cell) for cell in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
    headers = [header if header else f"__col_{idx}" for idx, header in enumerate(raw_headers, start=1)]
    rows: list[dict[str, Any]] = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        row = {headers[i]: values[i] for i in range(len(headers))}
        if any(clean_text(v) for v in row.values()):
            rows.append(row)
    return rows


def normalize_project_rows(rows: list[dict[str, Any]]) -> dict[str, Any]:
    projects: list[dict[str, Any]] = []
    sections: list[dict[str, Any]] = []
    warnings: list[str] = []

    current_project_name = ""
    current_project_short_name = ""

    for index, row in enumerate(rows, start=2):
        project_name = clean_text(find_value(row, ["工程全称", "项目全称", "工程名称"]))
        if project_name:
            current_project_name = project_name

        project_short_name = clean_text(find_value(row, ["工程简称", "项目简称"]))
        if project_short_name:
            current_project_short_name = project_short_name

        if not current_project_name:
            warnings.append(f"project row {index}: missing project name")
            continue

        project_record = {
            "project_name": current_project_name,
            "project_short_name": current_project_short_name or None,
            "project_alias": clean_text(find_value(row, ["别名"])) or None,
            "province": clean_text(find_value(row, ["省"])) or None,
            "city": clean_text(find_value(row, ["市"])) or None,
            "district": clean_text(find_value(row, ["区", "县"])) or None,
            "address_detail": clean_text(find_value(row, ["详细地址", "地址"])) or None,
            "geotech_soil": clean_text(find_value(row, ["地勘", "土质"])) or None,
            "machine_type": clean_text(find_value(row, ["机型"])) or None,
            "pipe_material": clean_text(find_value(row, ["管材"])) or None,
            "pipe_spec": clean_text(find_value(row, ["规格"])) or None,
            "legacy_note": clean_text(find_value(row, ["备注"])) or None,
        }
        if project_record not in projects:
            projects.append(project_record)

        section_code = clean_text(find_value(row, ["标段", "桩号"]))
        section_name = clean_text(find_value(row, ["标段名称"]))
        section_length = parse_number(find_value(row, ["长度", "里程", "米"]))
        design_slope = parse_number(find_value(row, ["设计坡度", "坡度"]))
        pipe_diameter_mm = parse_number(find_value(row, ["管径", "直径", "D"]))
        if section_code or section_name:
            sections.append(
                {
                    "project_name": current_project_name,
                    "section_code": section_code or None,
                    "section_name": section_name or None,
                    "section_length": section_length,
                    "design_slope": design_slope,
                    "pipe_diameter_mm": pipe_diameter_mm,
                }
            )

    return {"projects": projects, "sections": sections, "warnings": warnings}


def normalize_operator_rows(rows: list[dict[str, Any]]) -> dict[str, Any]:
    records: list[dict[str, Any]] = []
    warnings: list[str] = []
    last_project_name = ""
    last_section_text = ""
    last_log_date = ""

    for index, row in enumerate(rows, start=2):
        project_name = clean_text(find_value(row, ["工程名称", "项目名称"]))
        if project_name:
            last_project_name = project_name
        else:
            project_name = last_project_name

        section_text = clean_text(find_value(row, ["桩号", "标段"]))
        if section_text:
            last_section_text = section_text
        else:
            section_text = last_section_text

        log_date = clean_text(find_value(row, ["日期", "时间"]))
        if log_date:
            last_log_date = log_date
        else:
            log_date = last_log_date

        shift_raw = find_value(row, ["班次"])
        shift = normalize_shift(shift_raw)
        if shift is None and clean_text(shift_raw):
            warnings.append(f"operator row {index}: unknown shift {clean_text(shift_raw)}")

        deviation_combined = find_value(row, ["偏差"])
        deviation_direction, deviation_value_mm = split_deviation(deviation_combined)

        records.append(
            {
                "project_name": project_name or None,
                "section_text": section_text or None,
                "log_date": log_date or None,
                "shift": shift,
                "shift_cn": clean_text(shift_raw) or None,
                "operator_name": clean_text(find_value(row, ["操作员"])) or None,
                "completed_pipe_count": parse_number(find_value(row, ["本班", "下管"])) ,
                "cumulative_pipe_count": parse_number(find_value(row, ["累计", "下管"])) ,
                "footage_m": parse_number(find_value(row, ["顶进", "米"])) ,
                "main_thrust_mpa": parse_number(find_value(row, ["主推", "主顶力", "压力"])) ,
                "chamber_pressure_mpa": parse_number(find_value(row, ["舱压"])) ,
                "cutter_current_a": parse_number(find_value(row, ["刀盘", "电流"])) ,
                "screw_current_a": parse_number(find_value(row, ["螺旋", "电流"])) ,
                "soil_type": clean_text(find_value(row, ["土质"])) or None,
                "deviation_direction": deviation_direction,
                "deviation_value_mm": deviation_value_mm,
                "photo_keys_raw": clean_text(find_value(row, ["照片", "图片"])) or None,
                "remark": clean_text(find_value(row, ["备注"])) or None,
            }
        )

    return {"operator_logs": records, "warnings": warnings}


def normalize_manager_rows(rows: list[dict[str, Any]]) -> dict[str, Any]:
    records: list[dict[str, Any]] = []
    for row in rows:
        records.append(
            {
                "project_name": clean_text(find_value(row, ["工程名称", "项目名称"])) or None,
                "log_date": clean_text(find_value(row, ["日期"])) or None,
                "detail_html": clean_text(find_value(row, ["今日", "详情", "内容"])) or None,
                "tomorrow_plan": clean_text(find_value(row, ["明日", "计划"])) or None,
                "submitter_name": clean_text(find_value(row, ["提交人"])) or None,
                "submitted_at": clean_text(find_value(row, ["提交时间", "时间"])) or None,
            }
        )
    return {"manager_logs": records, "warnings": []}


def normalize_breakthrough_rows(rows: list[dict[str, Any]]) -> dict[str, Any]:
    records: list[dict[str, Any]] = []
    last_project_name = ""
    last_section_text = ""
    last_report_date = ""

    for row in rows:
        direction, deviation = split_deviation(find_value(row, ["偏差"]))
        project_name = clean_text(find_value(row, ["工程名称", "项目名称"]))
        if project_name:
            last_project_name = project_name
        else:
            project_name = last_project_name

        section_text = clean_text(find_value(row, ["标段", "桩号"]))
        if section_text:
            last_section_text = section_text
        else:
            section_text = last_section_text

        report_date = clean_text(find_value(row, ["贯通日期", "日期"]))
        if report_date:
            last_report_date = report_date
        else:
            report_date = last_report_date

        records.append(
            {
                "project_name": project_name or None,
                "section_text": section_text or None,
                "report_date": report_date or None,
                "start_date": clean_text(find_value(row, ["开工日期"])) or None,
                "jacking_days": parse_number(find_value(row, ["顶进天数", "天数"])) ,
                "pipe_diameter": clean_text(find_value(row, ["管径"])) or None,
                "total_length_m": parse_number(find_value(row, ["总长度", "长度"])) ,
                "total_pipe_count": parse_number(find_value(row, ["总下管", "管数量"])) ,
                "deviation_direction": direction,
                "deviation_value_mm": deviation,
                "other_notes": clean_text(find_value(row, ["其他事项", "备注"])) or None,
                "submitted_at": clean_text(find_value(row, ["提交时间"])) or None,
            }
        )
    return {"breakthrough_reports": records, "warnings": []}


def parse_legacy_excels(input_dir: Path) -> dict[str, Any]:
    files = {
        "project": "立项-工地及项目信息录入_*.xlsx",
        "operator_log": "操作员日志_*.xlsx",
        "manager_log": "管理员日志_*.xlsx",
        "breakthrough_report": "贯通报告_*.xlsx",
    }
    resolved: dict[str, Path] = {}
    for key, pattern in files.items():
        matched = sorted(input_dir.glob(pattern))
        if not matched:
            raise FileNotFoundError(f"missing file for {key}: {pattern}")
        resolved[key] = matched[-1]

    project_rows = read_rows(resolved["project"])
    operator_rows = read_rows(resolved["operator_log"])
    manager_rows = read_rows(resolved["manager_log"])
    breakthrough_rows = read_rows(resolved["breakthrough_report"])

    project_data = normalize_project_rows(project_rows)
    operator_data = normalize_operator_rows(operator_rows)
    manager_data = normalize_manager_rows(manager_rows)
    breakthrough_data = normalize_breakthrough_rows(breakthrough_rows)

    return {
        "generated_at": now_iso(),
        "input_dir": str(input_dir),
        "source_files": {k: str(v) for k, v in resolved.items()},
        "stats": {
            "project_rows": len(project_rows),
            "operator_rows": len(operator_rows),
            "manager_rows": len(manager_rows),
            "breakthrough_rows": len(breakthrough_rows),
            "project_records": len(project_data["projects"]),
            "section_records": len(project_data["sections"]),
            "operator_records": len(operator_data["operator_logs"]),
            "manager_records": len(manager_data["manager_logs"]),
            "breakthrough_records": len(breakthrough_data["breakthrough_reports"]),
        },
        "data": {
            "projects": project_data["projects"],
            "sections": project_data["sections"],
            "operator_logs": operator_data["operator_logs"],
            "manager_logs": manager_data["manager_logs"],
            "breakthrough_reports": breakthrough_data["breakthrough_reports"],
            "warnings": {
                "project": project_data["warnings"],
                "operator": operator_data["warnings"],
                "manager": manager_data["warnings"],
                "breakthrough": breakthrough_data["warnings"],
            },
        },
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Parse legacy Excel files into dry-run JSON output.")
    parser.add_argument(
        "--input-dir",
        default="doc/legacy_data",
        help="Directory containing legacy Excel files.",
    )
    parser.add_argument(
        "--output",
        default="task-runner/tmp/legacy_import_dry_run.json",
        help="Output JSON path.",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    input_dir = Path(args.input_dir)
    output_path = Path(args.output)

    if not input_dir.exists():
        raise FileNotFoundError(f"input directory not found: {input_dir}")

    result = parse_legacy_excels(input_dir)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"dry-run output written to: {output_path}")
    print(f"records: {json.dumps(result['stats'], ensure_ascii=False)}")


if __name__ == "__main__":
    main()
