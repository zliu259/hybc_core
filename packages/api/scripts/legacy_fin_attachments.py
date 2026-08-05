#!/usr/bin/env python3
"""附件映射器：把 mig_data.zip 分期附件回链到 legacy 合同/付款单据（迁移方案 §四）。

数据源：
- zip 内 对外付款申请_*（11 期）/ 对外付款合同结算管理_*（4 期）目录，
  每期含分期 Excel（付款「上传发票」「付款凭证上传」两列、合同「合同扫描件」列）
  与 Files/<字段>/ 实际文件（多附件记录被导出工具按「首文件名清洗后（去扩展名）」分组到子目录）。
- mig/ 汇总 Excel 提供范围过滤：合同只收流转完成的 560 个 CP 号；付款剔除测试行/缺陷行，
  row803 暂缓行单独产出（供财务确认后补导）。

产出（--output-dir，默认 mig/attachments/）：
- manifest.json / manifest_row803.json：legacy_data_id → 附件清单（r2_key 按方案 §四.3）
- upload_list.tsv / upload_list_row803.tsv：zip_member \t r2_key \t size
- unmatched.md：清洗后仍失配的引用、多义无法消歧的行、启发式配对复核项
- stats.json：各期匹配率、总文件数/字节、失配数、多义数
"""
from __future__ import annotations

import argparse
import datetime as dt
import io
import json
import re
import unicodedata
import zipfile
from collections import defaultdict, deque
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

PAYMENT_PREFIX = "对外付款申请_"
CONTRACT_PREFIX = "对外付款合同结算管理_"
PAYMENT_FIELDS = ("上传发票", "付款凭证上传")
CONTRACT_FIELDS = ("合同扫描件",)
# 展示名前缀：付款域区分发票/凭证，合同域用原名
FIELD_DISPLAY_PREFIX = {"上传发票": "发票/", "付款凭证上传": "凭证/", "合同扫描件": ""}

# 剔除/暂缓行（迁移方案 §八 决策 2/3，data_id 以汇总 Excel 实查冻结）
EXCLUDED_PAYMENT_DATA_IDS = {
    "606b0404cf6927000754299b",  # row2712 测试数据（收款方=测试，金额 99999）
    "6115ec6f8394d500077fbec5",  # row2496 EX20210800216 多列全空缺陷行
}
DEFERRED_PAYMENT_DATA_IDS = {
    "66c7f0a4103d1e9df86c0c9d",  # row803 EX20240801911 金额与大写不符，暂缓补导
}

# 导出工具同款清洗（经实际目录清单校准）：空格/&/+/#/:/=/,/% → _
# 注意匹配一律「精确优先、清洗兜底」——磁盘上存在含字面 % 的文件名，直接清洗会误伤。
CLEAN_RE = re.compile(r"[ &+#:=,%]")
# 导出工具的重名消歧后缀：子目录名尾部 / 文件名扩展名前的 -(n)
DUP_SUFFIX_RE = re.compile(r"-\((\d+)\)$")


def fix_zip_name(raw: str) -> str:
    """zip 条目名 cp437→utf-8 还原（UTF-8 标志位缺失时 zipfile 会错标）+ NFC 归一化。"""
    try:
        decoded = raw.encode("cp437").decode("utf-8")
    except (UnicodeEncodeError, UnicodeDecodeError):
        decoded = raw
    return unicodedata.normalize("NFC", decoded)


def clean_name(name: str) -> str:
    return CLEAN_RE.sub("_", name)


def split_ext(name: str) -> tuple[str, str]:
    dot = name.rfind(".")
    if dot <= 0:
        return name, ""
    return name[:dot], name[dot:]


def strip_dup_suffix(stem: str) -> str:
    return DUP_SUFFIX_RE.sub("", stem)


def dup_suffix_order(name: str) -> int:
    """重名消歧排序键：无后缀=0，-(n)=n（导出顺序与 Excel 行序一致）。"""
    stem, _ = split_ext(name)
    match = DUP_SUFFIX_RE.search(stem)
    return int(match.group(1)) if match else 0


def canon(name: str) -> str:
    """消歧比较用的规范形：清洗 + 去 -(n) 后缀 + 去导出工具补的前导下划线。"""
    stem, ext = split_ext(clean_name(name))
    return (strip_dup_suffix(stem) + ext).lstrip("_")


def as_text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text if text else None


@dataclass
class DiskFile:
    zip_member: str  # zipfile 原始条目名（解压时直接可用）
    period: str
    field: str
    subdir: str | None
    name: str
    size: int
    crc: int


@dataclass
class Record:
    period: str
    field: str
    data_id: str
    biz_no: str  # 流水号或合同号（人工清单定位用）
    row_no: int
    refs: list[str] = field(default_factory=list)


@dataclass
class Assignment:
    record: Record
    ref: str  # Excel 引用原文（作展示名）
    disk: DiskFile
    heuristic: bool = False  # 多义按行序启发式配对，需人工复核
    alternatives: list[str] = field(default_factory=list)


@dataclass
class Failure:
    record: Record
    ref: str
    kind: str  # unmatched=清洗后仍失配 / ambiguous=多义无法消歧
    detail: str = ""


class FieldMatcher:
    """单个 (期次, 字段) 命名空间内的记录→文件匹配与消歧。"""

    def __init__(self, period: str, field_name: str, files: list[DiskFile]) -> None:
        self.period = period
        self.field = field_name
        self.direct: dict[str, DiskFile] = {}
        self.subdirs: dict[str, list[DiskFile]] = defaultdict(list)
        for f in files:
            if f.subdir is None:
                self.direct[f.name] = f
            else:
                self.subdirs[f.subdir].append(f)
        for members in self.subdirs.values():
            members.sort(key=lambda f: (dup_suffix_order(f.name), f.name))
        # 引用名解析索引：精确名 + 清洗形 + 前导 _ 形
        self.name_lookup: dict[str, set[str]] = defaultdict(set)
        for name in list(self.direct) + [f.name for m in self.subdirs.values() for f in m]:
            for alias in {name, clean_name(name), clean_name(name).lstrip("_")}:
                self.name_lookup[alias].add(name)
        self.subdir_by_base: dict[str, list[str]] = defaultdict(list)
        for sub in self.subdirs:
            self.subdir_by_base[strip_dup_suffix(clean_name(sub)).lstrip("_")].append(sub)
        for subs in self.subdir_by_base.values():
            subs.sort(key=lambda s: (dup_suffix_order(s + ".x"), s))
        self.consumed_subdirs: set[str] = set()
        self.assignments: list[Assignment] = []
        self.failures: list[Failure] = []

    # ---- 引用拆分 ----

    def split_refs(self, cell: str) -> list[str]:
        """逗号贪心拆分：优先把相邻片段拼回能命中磁盘清单的最长文件名（防名内英文逗号）。"""
        parts = cell.split(",")
        refs: list[str] = []
        i = 0
        while i < len(parts):
            joined_end = i + 1
            for j in range(len(parts), i, -1):
                cand = ",".join(parts[i:j]).strip()
                if cand and (cand in self.name_lookup or clean_name(cand) in self.name_lookup):
                    joined_end = j
                    break
            ref = ",".join(parts[i:joined_end]).strip()
            if ref:
                refs.append(ref)
            i = joined_end
        return refs

    # ---- 匹配主流程 ----

    def match(self, records: list[Record]) -> None:
        remaining = self._match_subdir_groups(records)
        self._match_direct(remaining)

    def _subdir_candidates(self, rec: Record) -> list[str]:
        """子目录候选：子目录名 = 清洗后的首文件名（去扩展名），重名带 -(n) 后缀。"""
        seen: list[str] = []
        for ref in rec.refs:  # 以首引用为主，其余引用兜底
            base = strip_dup_suffix(split_ext(clean_name(ref))[0]).lstrip("_")
            for sub in self.subdir_by_base.get(base, []):
                if sub not in seen and self._compatible(sub, rec):
                    seen.append(sub)
            if seen:
                break
        return seen

    def _compatible(self, sub: str, rec: Record) -> bool:
        members = self.subdirs[sub]
        if len(members) != len(rec.refs):
            return False
        return sorted(canon(f.name) for f in members) == sorted(canon(r) for r in rec.refs)

    def _assign_subdir(self, rec: Record, sub: str, heuristic: bool, alts: list[str]) -> None:
        self.consumed_subdirs.add(sub)
        queues: dict[str, deque[DiskFile]] = defaultdict(deque)
        for f in self.subdirs[sub]:
            queues[canon(f.name)].append(f)
        for ref in rec.refs:
            disk = queues[canon(ref)].popleft()
            self.assignments.append(Assignment(rec, ref, disk, heuristic, alts))

    def _match_subdir_groups(self, records: list[Record]) -> list[Record]:
        """第一优先级：整记录 ↔ 子目录分组匹配；候选冲突用消耗法收敛，再按行序配对。"""
        pending: list[tuple[Record, list[str]]] = []
        remaining: list[Record] = []
        for rec in records:
            if len(rec.refs) < 2:  # 子目录只为多附件记录而建（实测无单文件子目录）
                remaining.append(rec)
                continue
            cands = self._subdir_candidates(rec)
            if cands:
                pending.append((rec, cands))
            else:
                remaining.append(rec)
        # 唯一候选先落位，逐轮释放其余记录的歧义
        progress = True
        while progress:
            progress = False
            for item in list(pending):
                rec, cands = item
                live = [c for c in cands if c not in self.consumed_subdirs]
                if len(live) == 1:
                    self._assign_subdir(rec, live[0], False, [])
                    pending.remove(item)
                    progress = True
        # 残余多义：按签名分组，行序 × 后缀序配对
        groups: dict[frozenset[str], list[Record]] = defaultdict(list)
        for rec, cands in pending:
            live = [c for c in cands if c not in self.consumed_subdirs]
            if not live:
                remaining.append(rec)  # 候选被别的记录消耗光，落到直挂匹配兜底
                continue
            groups[frozenset(live)].append(rec)
        for cand_set, recs in groups.items():
            cands = sorted(cand_set, key=lambda s: (dup_suffix_order(s + ".x"), s))
            recs.sort(key=lambda r: r.row_no)
            crc_sets = {tuple(sorted(f.crc for f in self.subdirs[c])) for c in cands}
            identical = len(crc_sets) == 1
            if identical or len(recs) == len(cands):
                for idx, rec in enumerate(recs):
                    sub = cands[idx % len(cands)]
                    self._assign_subdir(rec, sub, not identical, cands)
            else:
                for rec in recs:
                    self.failures.append(
                        Failure(rec, rec.refs[0], "ambiguous",
                                f"子目录候选 {len(cands)} 个内容不同且与记录数 {len(recs)} 不等: {cands}")
                    )
        return remaining

    def _match_direct(self, records: list[Record]) -> None:
        """第二优先级：直挂文件匹配；同名多文件按 crc 一致共享、行序启发式配对或进人工清单。"""
        occ_by_canon: dict[str, list[tuple[Record, str]]] = defaultdict(list)
        for rec in records:
            for ref in rec.refs:
                occ_by_canon[canon(ref)].append((rec, ref))
        for canon_key, occs in occ_by_canon.items():
            cands = sorted(
                (f for f in self.direct.values() if canon(f.name) == canon_key),
                key=lambda f: (dup_suffix_order(f.name), f.name),
            )
            if not cands:
                # 兜底：未被消耗子目录中的同名成员（唯一时才认）
                leftovers = [
                    f
                    for sub, members in self.subdirs.items()
                    if sub not in self.consumed_subdirs
                    for f in members
                    if canon(f.name) == canon_key
                ]
                if len(leftovers) == 1:
                    for rec, ref in occs:
                        self.assignments.append(Assignment(rec, ref, leftovers[0]))
                    continue
                kind = "ambiguous" if leftovers else "unmatched"
                detail = f"子目录残余同名 {len(leftovers)} 个" if leftovers else "清洗后仍无磁盘对应"
                for rec, ref in occs:
                    self.failures.append(Failure(rec, ref, kind, detail))
                continue
            if len(cands) == 1:
                for rec, ref in occs:  # 单文件可被多记录引用（内容唯一，无互覆风险）
                    self.assignments.append(Assignment(rec, ref, cands[0]))
                continue
            identical = len({f.crc for f in cands}) == 1
            if identical:
                for rec, ref in occs:
                    self.assignments.append(Assignment(rec, ref, cands[0]))
            elif len(occs) == len(cands):
                occs.sort(key=lambda t: t[0].row_no)
                names = [f.name for f in cands]
                for (rec, ref), disk in zip(occs, cands):
                    self.assignments.append(Assignment(rec, ref, disk, True, names))
            else:
                for rec, ref in occs:
                    self.failures.append(
                        Failure(rec, ref, "ambiguous",
                                f"直挂同名 {len(cands)} 个内容不同且与引用数 {len(occs)} 不等")
                    )


# ---- 数据装载 ----


def scan_zip(zf: zipfile.ZipFile) -> tuple[dict[tuple[str, str], list[DiskFile]], dict[str, str]]:
    """扫描 zip：返回 (期次,字段)→文件清单 与 期次→分期 Excel 条目名。"""
    files: dict[tuple[str, str], list[DiskFile]] = defaultdict(list)
    excels: dict[str, str] = {}
    for info in zf.infolist():
        decoded = fix_zip_name(info.filename)
        if decoded.startswith("__MACOSX") or info.is_dir():
            continue
        parts = decoded.split("/")
        if len(parts) < 3:
            continue
        period = parts[1]
        if not (period.startswith(PAYMENT_PREFIX) or period.startswith(CONTRACT_PREFIX)):
            continue
        if len(parts) == 3 and parts[2].endswith(".xlsx"):
            excels[period] = info.filename
            continue
        if len(parts) in (5, 6) and parts[2] == "Files":
            name = parts[-1]
            if name == ".DS_Store":
                continue
            subdir = parts[4] if len(parts) == 6 else None
            files[(period, parts[3])].append(
                DiskFile(info.filename, period, parts[3], subdir, name, info.file_size, info.CRC)
            )
    return files, excels


def load_period_rows(zf: zipfile.ZipFile, member: str) -> tuple[list[str], list[dict[str, Any]]]:
    wb = load_workbook(io.BytesIO(zf.read(member)), read_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    header = [as_text(h) or "" for h in next(rows_iter)]
    rows = []
    for row_no, row in enumerate(rows_iter, start=2):
        if row is None or all(v is None for v in row):
            continue
        width = min(len(header), len(row))  # read_only 模式下行元组可能短于表头
        rows.append({"_row_no": row_no, **{header[i]: row[i] for i in range(width)}})
    wb.close()
    return header, rows


def load_scope(contract_summary: Path, payment_summary: Path) -> tuple[set[str], dict[str, str]]:
    """汇总 Excel → 合同范围（流转完成 CP 号集合）与付款 data_id→流水号（核对用）。"""
    wb = load_workbook(contract_summary, read_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    header = [as_text(h) or "" for h in next(rows_iter)]
    idx = {h: i for i, h in enumerate(header)}
    cp_done: set[str] = set()
    for row in rows_iter:
        if row and as_text(row[idx["流程状态"]]) == "流转完成":
            cp = as_text(row[idx["合同号"]])
            if cp:
                cp_done.add(cp)
    wb.close()
    wb = load_workbook(payment_summary, read_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    header = [as_text(h) or "" for h in next(rows_iter)]
    idx = {h: i for i, h in enumerate(header)}
    payment_serials: dict[str, str] = {}
    for row in rows_iter:
        if not row:
            continue
        data_id = as_text(row[idx["data_id"]])
        if data_id:
            payment_serials[data_id] = as_text(row[idx["流水号"]]) or ""
    wb.close()
    return cp_done, payment_serials


# ---- 产出 ----


def period_token(period: str) -> str:
    """期次段：目录名去域前缀（如 2021Done / inprogress）。"""
    return period.split("_", 1)[1]


def sanitize_segment(seg: str) -> str:
    """R2 key 段清洗：字面 % 经 GET /object 双重解码必 404/400，统一改 _。"""
    return seg.replace("%", "_")


def build_row(assign: Assignment, domain: str) -> dict[str, Any]:
    disk = assign.disk
    safe_sub = sanitize_segment(disk.subdir) if disk.subdir else "_"
    safe_name = sanitize_segment(disk.name)
    r2_key = f"legacy/fin/{domain}/{period_token(disk.period)}/{disk.field}/{safe_sub}/{safe_name}"
    row: dict[str, Any] = {
        "zip_member": disk.zip_member,
        "r2_key": r2_key,
        "file_name": FIELD_DISPLAY_PREFIX[disk.field] + assign.ref,
        "size": disk.size,
    }
    if safe_name != disk.name or (disk.subdir and safe_sub != disk.subdir):
        row["renamed_from"] = (f"{disk.subdir}/" if disk.subdir else "") + disk.name
    return row


def write_outputs(out_dir: Path, manifest: dict, deferred: dict, failures: list[Failure],
                  reviews: list[Assignment], stats: dict) -> None:
    out_dir.mkdir(parents=True, exist_ok=True)
    (out_dir / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=1), encoding="utf-8")
    (out_dir / "manifest_row803.json").write_text(
        json.dumps(deferred, ensure_ascii=False, indent=1), encoding="utf-8")

    def upload_rows(source: dict) -> list[str]:
        seen: dict[str, str] = {}
        lines = []
        for entry in source.values():
            for row in entry["rows"]:
                key = row["r2_key"]
                if key in seen:
                    if seen[key] != row["zip_member"]:
                        raise SystemExit(f"r2_key 冲突（不同源文件同 key）: {key}")
                    continue
                seen[key] = row["zip_member"]
                lines.append(f"{row['zip_member']}\t{key}\t{row['size']}")
        return lines

    (out_dir / "upload_list.tsv").write_text("\n".join(upload_rows(manifest)) + "\n", "utf-8")
    (out_dir / "upload_list_row803.tsv").write_text(
        ("\n".join(upload_rows(deferred)) + "\n") if deferred else "", "utf-8")

    lines = ["# 附件映射人工复核清单", ""]
    lines.append("## 一、清洗后仍失配的引用（无磁盘对应）")
    lines.append("")
    lines.append("| 期次 | 字段 | data_id | 流水号/合同号 | 引用名 | 说明 |")
    lines.append("|---|---|---|---|---|---|")
    for f in failures:
        if f.kind == "unmatched":
            lines.append(f"| {f.record.period} | {f.record.field} | {f.record.data_id} "
                         f"| {f.record.biz_no} | {f.ref} | {f.detail} |")
    lines += ["", "## 二、多义无法消歧的行（未入 manifest）", "",
              "| 期次 | 字段 | data_id | 流水号/合同号 | 引用名 | 说明 |", "|---|---|---|---|---|---|"]
    for f in failures:
        if f.kind == "ambiguous":
            lines.append(f"| {f.record.period} | {f.record.field} | {f.record.data_id} "
                         f"| {f.record.biz_no} | {f.ref} | {f.detail} |")
    lines += ["", "## 三、启发式配对（已入 manifest，同名不同内容按行序配对，请抽查）", "",
              "| 期次 | 字段 | data_id | 流水号/合同号 | 引用名 | 配到 | 候选 |", "|---|---|---|---|---|---|---|"]
    for a in reviews:
        target = (f"{a.disk.subdir}/" if a.disk.subdir else "") + a.disk.name
        lines.append(f"| {a.record.period} | {a.record.field} | {a.record.data_id} "
                     f"| {a.record.biz_no} | {a.ref} | {target} | {'; '.join(a.alternatives)} |")
    (out_dir / "unmatched.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    (out_dir / "stats.json").write_text(
        json.dumps(stats, ensure_ascii=False, indent=1), encoding="utf-8")


# ---- 主流程 ----


def main() -> None:
    repo_root = Path(__file__).resolve().parents[3]
    parser = argparse.ArgumentParser(description="legacy 财务附件映射器（迁移方案 §四）")
    parser.add_argument("--zip", default="/Users/zihaoliu/backup/mig_data.zip",
                        help="mig_data.zip 路径（只读）")
    parser.add_argument("--contract-summary",
                        default=str(repo_root / "mig/对外付款合同结算管理_20260310214755.xlsx"))
    parser.add_argument("--payment-summary",
                        default=str(repo_root / "mig/对外付款申请_20260310214713.xlsx"))
    parser.add_argument("--output-dir", default=str(repo_root / "mig/attachments"))
    args = parser.parse_args()

    cp_done, payment_serials = load_scope(Path(args.contract_summary), Path(args.payment_summary))
    if len(cp_done) != 560:
        raise SystemExit(f"合同范围异常：流转完成 CP 号 {len(cp_done)} 个，预期 560")

    zf = zipfile.ZipFile(args.zip)
    disk_files, excels = scan_zip(zf)

    manifest: dict[str, dict[str, Any]] = {}
    deferred: dict[str, dict[str, Any]] = {}
    failures: list[Failure] = []
    reviews: list[Assignment] = []
    stats_periods: dict[str, dict[str, Any]] = {}
    skipped = {"contract_out_of_scope": 0, "payment_excluded": 0, "payment_deferred": 0}

    for period in sorted(excels):
        is_contract = period.startswith(CONTRACT_PREFIX)
        domain = "contract" if is_contract else "payment"
        fields = CONTRACT_FIELDS if is_contract else PAYMENT_FIELDS
        biz_col = "合同号" if is_contract else "流水号"
        _, rows = load_period_rows(zf, excels[period])

        disk_in_period = [f for fname in fields for f in disk_files.get((period, fname), [])]
        pstat = {"rows": len(rows), "rows_in_scope": 0, "refs": 0, "matched": 0,
                 "unmatched": 0, "ambiguous": 0, "heuristic_review": 0,
                 "files": 0, "bytes": 0,
                 "disk_files": len(disk_in_period),
                 "disk_bytes": sum(f.size for f in disk_in_period)}
        seen_keys: set[str] = set()
        matchers = {fname: FieldMatcher(period, fname, disk_files.get((period, fname), []))
                    for fname in fields}
        records: dict[str, list[Record]] = {fname: [] for fname in fields}
        row_meta: dict[str, str] = {}

        for row in rows:
            data_id = as_text(row.get("data_id"))
            if not data_id:
                continue
            biz_no = as_text(row.get(biz_col)) or ""
            if is_contract:
                if biz_no not in cp_done:  # 15 条放弃合同 + CP26031/26032 一律跳过
                    skipped["contract_out_of_scope"] += 1
                    continue
            else:
                if data_id in EXCLUDED_PAYMENT_DATA_IDS:
                    skipped["payment_excluded"] += 1
                    continue
                if data_id in DEFERRED_PAYMENT_DATA_IDS:
                    skipped["payment_deferred"] += 1
            pstat["rows_in_scope"] += 1
            row_meta[data_id] = biz_no
            for fname in fields:
                cell = as_text(row.get(fname))
                if not cell:
                    continue
                cell = unicodedata.normalize("NFC", cell)
                rec = Record(period, fname, data_id, biz_no, row["_row_no"])
                rec.refs = matchers[fname].split_refs(cell)
                if rec.refs:
                    records[fname].append(rec)
                    pstat["refs"] += len(rec.refs)

        for fname in fields:
            matcher = matchers[fname]
            matcher.match(records[fname])
            for assign in matcher.assignments:
                target = deferred if assign.record.data_id in DEFERRED_PAYMENT_DATA_IDS else manifest
                entry = target.setdefault(
                    assign.record.data_id,
                    {"domain": domain, "biz_no": row_meta[assign.record.data_id],
                     "period": period_token(period), "rows": []},
                )
                row = build_row(assign, domain)
                pstat["matched"] += 1
                # 同一单元格把同名引用写两次而磁盘只有一份（内容相同被导出工具合并）时去重
                if any(r["r2_key"] == row["r2_key"] and r["file_name"] == row["file_name"]
                       for r in entry["rows"]):
                    continue
                entry["rows"].append(row)
                if row["r2_key"] not in seen_keys:
                    seen_keys.add(row["r2_key"])
                    pstat["files"] += 1
                    pstat["bytes"] += assign.disk.size
                if assign.heuristic:
                    reviews.append(assign)
                    pstat["heuristic_review"] += 1
            for fail in matcher.failures:
                failures.append(fail)
                pstat["unmatched" if fail.kind == "unmatched" else "ambiguous"] += 1

        pstat["match_rate"] = round(pstat["matched"] / pstat["refs"], 4) if pstat["refs"] else 1.0
        stats_periods[period] = pstat

    total = {k: sum(p[k] for p in stats_periods.values())
             for k in ("rows", "rows_in_scope", "refs", "matched", "unmatched",
                       "ambiguous", "heuristic_review", "files", "bytes",
                       "disk_files", "disk_bytes")}
    total["match_rate"] = round(total["matched"] / total["refs"], 4) if total["refs"] else 1.0
    # 未被任何在范围记录引用的磁盘文件（含范围外行的附件），供上传对账时解释差额
    total["disk_files_unreferenced"] = total["disk_files"] - total["files"]
    stats = {
        "generated_at": dt.datetime.now().replace(microsecond=0).isoformat(),
        "zip": args.zip,
        "scope": {"contract_cp_done": len(cp_done), "skipped": skipped,
                  "manifest_data_ids": len(manifest), "deferred_data_ids": len(deferred)},
        "total": total,
        "periods": stats_periods,
    }
    write_outputs(Path(args.output_dir), manifest, deferred, failures, reviews, stats)
    print(json.dumps({"total": total, "scope": stats["scope"]}, ensure_ascii=False, indent=1))


if __name__ == "__main__":
    main()
