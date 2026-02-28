#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import sqlite3
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def as_text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text if text else None


def load_rows(path: Path) -> list[dict[str, Any]]:
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    header_row = 1
    best = 0
    for idx in range(1, min(ws.max_row, 20) + 1):
        values = [as_text(ws.cell(idx, col).value) for col in range(1, ws.max_column + 1)]
        count = sum(1 for x in values if x)
        if count > best:
            best = count
            header_row = idx
    headers = [as_text(ws.cell(header_row, col).value) or "" for col in range(1, ws.max_column + 1)]
    rows: list[dict[str, Any]] = []
    for ridx in range(header_row + 1, ws.max_row + 1):
        values = [ws.cell(ridx, col).value for col in range(1, len(headers) + 1)]
        if not any(as_text(v) for v in values):
            continue
        rows.append({headers[i]: values[i] for i in range(len(headers)) if headers[i]})
    return rows


def discover_local_db_path(repo_root: Path) -> Path | None:
    base = repo_root / "packages" / "api" / ".wrangler" / "state" / "v3" / "d1" / "miniflare-D1DatabaseObject"
    if not base.exists():
        return None
    candidates = sorted(base.glob("*.sqlite"), key=lambda p: p.stat().st_mtime, reverse=True)
    return candidates[0] if candidates else None


def unique_text(values: list[Any]) -> list[str]:
    result: list[str] = []
    seen: set[str] = set()
    for value in values:
        text = as_text(value)
        if not text or text in seen:
            continue
        seen.add(text)
        result.append(text)
    return result


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate initial mapping file for legacy finance import")
    parser.add_argument("--repo-root", default=str(Path(__file__).resolve().parents[3]))
    parser.add_argument("--input-dir", default="doc/legacy_data/fin_data")
    parser.add_argument("--output", default="doc/legacy_data/fin_data_mappings.generated.json")
    parser.add_argument("--db-path", default="")
    args = parser.parse_args()

    repo_root = Path(args.repo_root).resolve()
    input_dir = (repo_root / args.input_dir).resolve()
    output = (repo_root / args.output).resolve()
    db_path = Path(args.db_path).resolve() if args.db_path else discover_local_db_path(repo_root)
    if db_path is None or not db_path.exists():
        raise RuntimeError("Local D1 sqlite not found. Run migrations first or pass --db-path.")

    conn = sqlite3.connect(db_path)

    project_rows = conn.execute("SELECT id, project_code, project_name FROM project ORDER BY id DESC").fetchall()
    entity_rows = conn.execute("SELECT id, entity_name FROM entity").fetchall()

    project_by_code = {str(row[1]): int(row[0]) for row in project_rows if row[1]}
    project_by_name = {str(row[2]): int(row[0]) for row in project_rows if row[2]}
    entity_by_name = {str(row[1]): int(row[0]) for row in entity_rows if row[1]}

    contract_rows = load_rows(input_dir / "对外付款合同结算管理_20260228101109.xlsx")
    payment_rows = load_rows(input_dir / "对外付款申请_20260228101021.xlsx")
    envelope_rows = load_rows(input_dir / "个人报销预申请_20260228101235.xlsx")
    batch_rows = load_rows(input_dir / "个人报销集中申报_20260228101252.xlsx")

    project_candidates = unique_text(
        [
            *[row.get("所属工程") for row in contract_rows],
            *[row.get("工程全称") for row in contract_rows],
            *[row.get("所属项目/工程") for row in payment_rows],
            *[row.get("工程全称") for row in payment_rows],
            *[row.get("所属工程") for row in envelope_rows],
            *[row.get("所属项目/工程") for row in batch_rows],
        ]
    )

    user_candidates = unique_text(
        [
            *[row.get("经办人") for row in contract_rows],
            *[row.get("提交人") for row in contract_rows],
            *[row.get("申请人") for row in payment_rows],
            *[row.get("办理人") for row in payment_rows],
            *[row.get("提交人") for row in payment_rows],
            *[row.get("提交人") for row in envelope_rows],
            *[row.get("待报销人") for row in batch_rows],
            *[row.get("提交人") for row in batch_rows],
        ]
    )

    entity_candidates = unique_text([row.get("合同乙方") for row in contract_rows])

    project_map: dict[str, int | None] = {}
    for name in project_candidates:
        if name in project_by_code:
            project_map[name] = project_by_code[name]
        elif name in project_by_name:
            project_map[name] = project_by_name[name]
        else:
            project_map[name] = None

    entity_map: dict[str, int | None] = {}
    for name in entity_candidates:
        entity_map[name] = entity_by_name.get(name)

    user_map: dict[str, str | None] = {name: None for name in user_candidates}

    payload = {
        "default_userid": "legacy_import",
        "default_project_id": int(project_rows[0][0]) if project_rows else None,
        "auto_create_entity": True,
        "project_map": project_map,
        "user_map": user_map,
        "entity_map": entity_map,
        "unresolved": {
            "project_names": [k for k, v in project_map.items() if v is None],
            "user_names": [k for k, v in user_map.items() if v is None],
            "entity_names": [k for k, v in entity_map.items() if v is None],
        },
    }

    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    print(
        json.dumps(
            {
                "ok": True,
                "output": str(output),
                "project_total": len(project_map),
                "project_resolved": len([1 for v in project_map.values() if v is not None]),
                "user_total": len(user_map),
                "entity_total": len(entity_map),
                "entity_resolved": len([1 for v in entity_map.values() if v is not None]),
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
