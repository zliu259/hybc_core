#!/usr/bin/env python3
"""旧财务付款申请补迁：生产幂等 SQL 批文件生成器（迁移方案 v1.2 §3.1/§3.4/第 1 步）。

读取 mig/对外付款申请_*.xlsx + mig/fin_prod_mappings.json + mig/attachments/manifest*.json，
生成 mig/sql/ 下全部批文件（含 rollback/ 与 verify.sql）。只生成文件，不连任何数据库。

冻结口径（勿改）：
- 剔除 row2712(测试, 606b0404cf6927000754299b)、row2496(EX20210800216, 6115ec6f8394d500077fbec5)；
- row803(EX20240801911, 66c7f0a4103d1e9df86c0c9d) 暂缓，单独产出 90_row803_deferred.sql；
- 本批 2708 行 = paid 2622 + cancelled 86；
- hex/未映射用户一律落 fallback_userid（legacy-import）；
- category/contract_type 原文照抄；payment_method 现金→现金、其余→对公转账；
- amount_cn 用 services/amount-cn.ts 同款算法重算（含 node 一致性烟测）。

映射文件更新后重跑本脚本即可重新生成全部产物。
"""
from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import json
import math
import random
import re
import subprocess
import tempfile
import unicodedata
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

# ---------- 冻结常量 ----------
EXCLUDED_DATA_IDS = {
    "606b0404cf6927000754299b": "row2712 测试数据（收款方=测试，金额 99999）",
    "6115ec6f8394d500077fbec5": "row2496 EX20210800216 多列全空缺陷行",
}
DEFERRED_DATA_ID = "66c7f0a4103d1e9df86c0c9d"  # row803 EX20240801911，金额待财务对流水
ORPHAN_CONTRACT_NO = "C22002"  # 付款表无主合同号 → contract_id NULL
PAYEE_LEGACY_SOURCE = "fin_data_auto_20260805"
BATCH_LIMIT = 500
EXPECT_TOTAL, EXPECT_PAID, EXPECT_CANCELLED, EXPECT_CONTRACT_LINKED = 2708, 2622, 86, 964

MIG_NOTE_PREFIX = "【迁移】旧系统流程状态："

# ---------- 基础转换 ----------

def as_text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    return text if text else None


def as_amount(value: Any) -> float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = as_text(value)
    if not text:
        return None
    try:
        return float(text.replace(",", "").replace("，", ""))
    except ValueError:
        return None


def as_iso_datetime(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, dt.datetime):
        return value.replace(microsecond=0).isoformat() + "Z"
    if isinstance(value, dt.date):
        return dt.datetime(value.year, value.month, value.day).isoformat() + "Z"
    text = as_text(value)
    if not text:
        return None
    normalized = text.replace("/", "-")
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d"):
        try:
            return dt.datetime.strptime(normalized, fmt).isoformat() + "Z"
        except ValueError:
            continue
    return text


def as_iso_date(value: Any) -> str | None:
    if isinstance(value, dt.datetime) or isinstance(value, dt.date):
        return value.strftime("%Y-%m-%d")
    text = as_text(value)
    if not text:
        return None
    normalized = text.replace("/", "-")
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d"):
        try:
            return dt.datetime.strptime(normalized, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return text


def norm_ws(value: str | None) -> str:
    """去半角空格/全角空格/制表符——与 DB 侧 REPLACE 链严格同构。"""
    if not value:
        return ""
    return value.replace(" ", "").replace("　", "").replace("\t", "")


def nfkc_key(value: str | None) -> str:
    """项目匹配用：NFKC + 去所有空白。"""
    if not value:
        return ""
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", value))


NORM_SQL = "REPLACE(REPLACE(REPLACE({col}, ' ', ''), '　', ''), char(9), '')"


def sql_str(value: str | None) -> str:
    """SQL 字面量：单引号翻倍；换行转 char(10) 拼接，保证一行一句。"""
    if value is None:
        return "NULL"
    text = value.replace("\r\n", "\n").replace("\r", "\n").replace("'", "''")
    if "\n" not in text:
        return f"'{text}'"
    parts = text.split("\n")
    return "||char(10)||".join(f"'{p}'" for p in parts)


def sql_num(value: float) -> str:
    rounded = round(value + 0.0, 2)
    if rounded == int(rounded):
        return str(int(rounded))
    return f"{rounded:.2f}"


# ---------- amount-cn.ts 的 Python 移植（与 JS 逐语义对齐） ----------
CN_DIGITS = ["零", "壹", "贰", "叁", "肆", "伍", "陆", "柒", "捌", "玖"]
CN_UNITS = ["", "拾", "佰", "仟"]
CN_SECTION_UNITS = ["", "万", "亿", "万亿"]


def _convert_section(section: int) -> str:
    result = ""
    unit_index = 0
    need_zero = False
    while section > 0:
        digit = section % 10
        if digit == 0:
            need_zero = True
        else:
            if need_zero:
                result = CN_DIGITS[0] + result
                need_zero = False
            result = CN_DIGITS[digit] + CN_UNITS[unit_index] + result
        section //= 10
        unit_index += 1
    return result


def _convert_integer(n: int) -> str:
    if n == 0:
        return CN_DIGITS[0]
    result = ""
    section_index = 0
    need_zero = False
    while n > 0:
        section = n % 10000
        if section != 0:
            section_str = _convert_section(section)
            if need_zero:
                result = section_str + CN_SECTION_UNITS[section_index] + "零" + result
            else:
                result = section_str + CN_SECTION_UNITS[section_index] + result
            need_zero = False
        else:
            need_zero = True
        n //= 10000
        section_index += 1
    return result


def to_amount_chinese(amount: float) -> str:
    if amount < 0:
        return "负" + to_amount_chinese(-amount)
    if amount == 0:
        return "零元整"
    rounded = math.floor(amount * 100 + 0.5)  # JS Math.round：half 向上
    integer_part = rounded // 100
    decimal_part = rounded % 100
    jiao = decimal_part // 10
    fen = decimal_part % 10
    result = ""
    if integer_part > 0:
        result = _convert_integer(integer_part) + "元"
    if jiao == 0 and fen == 0:
        result += "整"
    else:
        if integer_part > 0 and jiao == 0:
            result += "零"
        if jiao > 0:
            result += CN_DIGITS[jiao] + "角"
        if fen > 0:
            result += CN_DIGITS[fen] + "分"
    return result


def run_amount_cn_parity(ts_path: Path, amounts: list[float]) -> dict[str, Any]:
    """node 一致性烟测：把 amount-cn.ts 剥离类型注解后原样在 node 里执行比对。"""
    ts_src = ts_path.read_text(encoding="utf-8")
    js_src = ts_src.replace("export const", "const")
    js_src = re.sub(r"\(([A-Za-z_]+): *[A-Za-z]+\): *[A-Za-z]+ *=>", r"(\1) =>", js_src)
    sample = [0, 0.01, 0.1, 0.5, 1, 10.05, 100.2, -1012, -0.07, 99999, 250000, 9909.63,
              100000001.01, 1000000000000.34, 35200, 19532.03]
    rng = random.Random(20260805)
    randoms = [round(rng.uniform(-1e6, 1e8), 2) for _ in range(200 - len(sample))]
    test_values = sample + randoms + sorted({round(a, 2) for a in amounts})
    with tempfile.TemporaryDirectory() as tmp:
        in_file = Path(tmp) / "amounts.json"
        js_file = Path(tmp) / "amount_cn_parity.mjs"
        in_file.write_text(json.dumps(test_values), encoding="utf-8")
        js_file.write_text(
            js_src
            + "\nimport { readFileSync } from 'node:fs';\n"
            + f"const values = JSON.parse(readFileSync('{in_file}', 'utf-8'));\n"
            + "console.log(JSON.stringify(values.map(v => toAmountChinese(v))));\n",
            encoding="utf-8",
        )
        proc = subprocess.run(["node", str(js_file)], capture_output=True, text=True, check=True)
    js_out = json.loads(proc.stdout)
    py_out = [to_amount_chinese(v) for v in test_values]
    mismatches = [
        {"amount": test_values[i], "python": py_out[i], "node": js_out[i]}
        for i in range(len(test_values))
        if py_out[i] != js_out[i]
    ]
    return {"total_cases": len(test_values), "mismatches": mismatches, "passed": not mismatches}


# ---------- 加载 ----------

def load_rows(excel_path: Path) -> list[dict[str, Any]]:
    wb = load_workbook(excel_path, data_only=True)
    ws = wb.active
    header = [as_text(c.value) or "" for c in ws[1]]
    rows: list[dict[str, Any]] = []
    for row_idx in range(2, ws.max_row + 1):
        values = [ws.cell(row_idx, col).value for col in range(1, len(header) + 1)]
        if not any(as_text(v) is not None for v in values):
            continue
        rows.append({"__row_no": row_idx, **{header[i]: values[i] for i in range(len(header)) if header[i]}})
    return rows


def load_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


# ---------- 解析器 ----------
class Resolver:
    def __init__(self, mappings: dict[str, Any]):
        self.fallback_userid: str = str(mappings.get("fallback_userid") or "legacy-import")
        self.user_map: dict[str, str] = {}
        for name, item in (mappings.get("user_map") or {}).items():
            userid = item.get("userid") if isinstance(item, dict) else item
            if userid:
                self.user_map[name] = str(userid)
        self.project_map: dict[str, int] = {}
        self.project_map_nfkc: dict[str, int] = {}
        for name, item in (mappings.get("project_map") or {}).items():
            pid = item.get("project_id") if isinstance(item, dict) else item
            if pid is None:
                continue
            self.project_map[name] = int(pid)
            self.project_map_nfkc.setdefault(nfkc_key(name), int(pid))
        self.entity_map: dict[str, int] = {}
        for name, item in (mappings.get("entity_map") or {}).items():
            eid = item.get("entity_id") if isinstance(item, dict) else item
            if eid is not None:
                self.entity_map[name] = int(eid)
        self.stats = {"user_mapped": 0, "user_fallback": 0, "project_exact": 0,
                      "project_nfkc": 0, "project_miss": 0, "entity_hit": 0, "entity_miss": 0}

    def userid(self, name: str | None) -> str | None:
        if not name:
            return None
        if name in self.user_map:
            self.stats["user_mapped"] += 1
            return self.user_map[name]
        self.stats["user_fallback"] += 1
        return self.fallback_userid

    def project_id(self, name: str | None) -> int | None:
        if not name:
            return None
        if name in self.project_map:
            self.stats["project_exact"] += 1
            return self.project_map[name]
        key = nfkc_key(name)
        if key in self.project_map_nfkc:
            self.stats["project_nfkc"] += 1
            return self.project_map_nfkc[key]
        self.stats["project_miss"] += 1
        return None

    def entity_id(self, pay_method_raw: str | None) -> int | None:
        if not pay_method_raw:
            return None
        if pay_method_raw in self.entity_map:
            self.stats["entity_hit"] += 1
            return self.entity_map[pay_method_raw]
        self.stats["entity_miss"] += 1
        return None


# ---------- 行构建 ----------

def build_attachments_json(manifest_entry: dict[str, Any] | None) -> str:
    if not manifest_entry:
        return "[]"
    items = [
        {"file_key": r["r2_key"], "file_name": r["file_name"], "file_size": r["size"]}
        for r in manifest_entry.get("rows", [])
    ]
    return json.dumps(items, ensure_ascii=False, separators=(",", ":"))


def payment_insert_sql(
    row: dict[str, Any],
    resolver: Resolver,
    manifest_entry: dict[str, Any] | None,
    voucher_raw: str | None,
    prevalidation: list[dict[str, Any]],
    warnings: list[str],
) -> tuple[str, dict[str, Any]]:
    data_id = as_text(row.get("data_id"))
    assert data_id
    serial_no = as_text(row.get("流水号"))
    if not serial_no:
        serial_no = f"EX-LEGACY-{data_id[:8]}"
        warnings.append(f"data_id={data_id}（行 {row['__row_no']}）流水号为空，合成 {serial_no}")

    flow_status = as_text(row.get("流程状态")) or ""
    status = "paid" if flow_status == "流转完成" else "cancelled"

    project_name = as_text(row.get("所属项目/工程"))
    project_id = resolver.project_id(project_name)
    applicant_name = as_text(row.get("申请人"))
    applicant_userid = resolver.userid(applicant_name)
    handler_name = as_text(row.get("办理人"))
    paid_by_userid = resolver.userid(handler_name)

    payee_name = as_text(row.get("收款单位/收款人"))
    payee_bank = as_text(row.get("收款方开户行信息"))
    payee_account = as_text(row.get("收款方账号"))
    payee_phone = as_text(row.get("收款方电话"))

    contract_no = as_text(row.get("合同号"))
    amount = as_amount(row.get("付款金额"))
    expected_date = as_iso_date(row.get("预计付款日期"))
    category = as_text(row.get("报销类别"))
    content = as_text(row.get("付款内容（详情）"))
    pay_method_raw = as_text(row.get("付款方式"))
    payment_method = "现金" if pay_method_raw == "现金" else "对公转账"
    payment_entity_id = resolver.entity_id(pay_method_raw)
    created_at = as_iso_datetime(row.get("提交时间"))
    updated_at = as_iso_datetime(row.get("更新时间"))
    paid_at = as_iso_datetime(row.get("实际付款日期")) if status == "paid" else None
    if status == "paid" and not paid_at:
        paid_at = updated_at
        warnings.append(f"data_id={data_id} 流转完成但实际付款日期为空，paid_at 回退更新时间")

    remark = as_text(row.get("备注"))
    if status == "cancelled":
        note = f"{MIG_NOTE_PREFIX}{flow_status}"
        node = as_text(row.get("当前节点"))
        if node:
            note += f"（当前节点：{node}）"
        remark = f"{remark}\n{note}" if remark else note

    # NOT NULL 预校验（目标表 NOT NULL 列）
    missing = [
        col
        for col, val in [
            ("serial_no", serial_no),
            ("project_id", project_id),
            ("applicant_userid", applicant_userid),
            ("expected_payment_date", expected_date),
            ("category", category),
            ("payee_id(收款单位/收款人)", payee_name),
            ("amount", amount),
            ("content", content),
            ("created_at", created_at),
            ("updated_at", updated_at),
        ]
        if val is None
    ]
    if missing:
        prevalidation.append({"row_no": row["__row_no"], "data_id": data_id,
                              "serial_no": serial_no, "missing": missing})

    amount_cn = to_amount_chinese(round(amount, 2)) if amount is not None else None

    if contract_no and contract_no != ORPHAN_CONTRACT_NO:
        contract_expr = f"(SELECT id FROM contract WHERE contract_no={sql_str(contract_no)} LIMIT 1)"
    else:
        contract_expr = "NULL"
    payee_expr = (
        "(SELECT id FROM payee WHERE "
        + NORM_SQL.format(col="name") + f"={sql_str(norm_ws(payee_name))} AND "
        + NORM_SQL.format(col="bank_account") + f"={sql_str(norm_ws(payee_account))}"
        + " ORDER BY id LIMIT 1)"
    )

    cols = (
        "serial_no, contract_id, project_id, applicant_userid, expected_payment_date, category, "
        "payee_id, amount, amount_cn, content, payment_method, status, approval_instance_id, "
        "ai_risk_status, attachments_json, remark, created_at, updated_at, paid_at, paid_by_userid, "
        "payment_account, payment_entity_id, legacy_data_id, legacy_applicant_name, "
        "legacy_project_name, legacy_project_full_name, legacy_contract_no, legacy_contract_status, "
        "legacy_invoice_type, legacy_payee_name, legacy_payee_bank_name, legacy_payee_account, "
        "legacy_payee_phone, legacy_handler_name, legacy_specified_pay_flag, "
        "legacy_specified_pay_date, legacy_voucher_keys_raw, legacy_flow_status, "
        "legacy_current_node, legacy_current_owner, legacy_submitter_name, legacy_submitted_at"
    )
    vals = ", ".join([
        sql_str(serial_no),
        contract_expr,
        str(project_id) if project_id is not None else "NULL",
        sql_str(applicant_userid),
        sql_str(expected_date),
        sql_str(category),
        payee_expr,
        sql_num(amount) if amount is not None else "NULL",
        sql_str(amount_cn),
        sql_str(content),
        sql_str(payment_method),
        sql_str(status),
        "NULL",
        "'none'",
        sql_str(build_attachments_json(manifest_entry)),
        sql_str(remark),
        sql_str(created_at),
        sql_str(updated_at),
        sql_str(paid_at),
        sql_str(paid_by_userid),
        sql_str(pay_method_raw),
        str(payment_entity_id) if payment_entity_id is not None else "NULL",
        sql_str(data_id),
        sql_str(applicant_name),
        sql_str(project_name),
        sql_str(as_text(row.get("工程全称"))),
        sql_str(contract_no),
        sql_str(as_text(row.get("合同状态"))),
        sql_str(as_text(row.get("发票类型"))),
        sql_str(payee_name),
        sql_str(payee_bank),
        sql_str(payee_account),
        sql_str(payee_phone),
        sql_str(handler_name),
        sql_str(as_text(row.get("指定日期付款"))),
        sql_str(as_iso_date(row.get("指定付款日期"))),
        sql_str(voucher_raw),
        sql_str(flow_status),
        sql_str(as_text(row.get("当前节点"))),
        sql_str(as_text(row.get("当前负责人"))),
        sql_str(as_text(row.get("提交人"))),
        sql_str(as_iso_datetime(row.get("提交时间"))),
    ])
    stmt = (
        f"INSERT INTO payment_request ({cols}) SELECT {vals} "
        f"WHERE NOT EXISTS (SELECT 1 FROM payment_request WHERE legacy_data_id={sql_str(data_id)});"
    )
    meta = {
        "data_id": data_id,
        "status": status,
        "amount": round(amount, 2) if amount is not None else None,
        "contract_no": contract_no,
        "contract_subquery": contract_expr != "NULL",
        "year": (paid_at or created_at or "")[:4],
    }
    return stmt, meta


def payee_insert_sql(first_row: dict[str, Any], resolver: Resolver) -> str:
    name = as_text(first_row.get("收款单位/收款人"))
    bank_name = as_text(first_row.get("收款方开户行信息")) or ""
    account = norm_ws(as_text(first_row.get("收款方账号")))
    phone = as_text(first_row.get("收款方电话"))
    submitted_by = as_text(first_row.get("提交人"))
    submitted_at = as_iso_datetime(first_row.get("提交时间"))
    updated_at = as_iso_datetime(first_row.get("更新时间")) or submitted_at
    predicate = (
        NORM_SQL.format(col="name") + f"={sql_str(norm_ws(name))} AND "
        + NORM_SQL.format(col="bank_account") + f"={sql_str(account)}"
    )
    vals = ", ".join([
        sql_str(name),
        sql_str(bank_name),
        sql_str(account),
        sql_str(phone),
        "'external'",
        "'active'",
        sql_str(resolver.fallback_userid),
        sql_str(submitted_at),
        sql_str(updated_at),
        "NULL",
        sql_str(PAYEE_LEGACY_SOURCE),
        sql_str(submitted_by),
        sql_str(submitted_at),
    ])
    return (
        "INSERT INTO payee (name, bank_name, bank_account, phone, type, status, created_by_userid, "
        "created_at, updated_at, legacy_data_id, legacy_source, legacy_submitted_by, legacy_submitted_at) "
        f"SELECT {vals} WHERE NOT EXISTS (SELECT 1 FROM payee WHERE {predicate});"
    )


# ---------- 文件写出 ----------

def write_sql(path: Path, header_lines: list[str], statements: list[str]) -> int:
    lines = [f"-- {h}" for h in header_lines] + statements
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return len(statements)


def write_batched(dir_: Path, stem: str, header: str, statements: list[str], counts: dict[str, int]) -> list[str]:
    """≤500 语句一批；单批时不带序号后缀。返回文件名清单。"""
    names: list[str] = []
    if len(statements) <= BATCH_LIMIT:
        name = f"{stem}.sql"
        counts[name] = write_sql(dir_ / name, [header], statements)
        return [name]
    for i in range(0, len(statements), BATCH_LIMIT):
        name = f"{stem}_{i // BATCH_LIMIT + 1:03d}.sql"
        counts[name] = write_sql(dir_ / name, [header + f"（第 {i // BATCH_LIMIT + 1} 批）"], statements[i:i + BATCH_LIMIT])
        names.append(name)
    return names


# ---------- 主流程 ----------

def main() -> None:
    parser = argparse.ArgumentParser(description="Emit idempotent remote SQL for legacy fin payment migration")
    repo_root_default = Path(__file__).resolve().parents[3]
    parser.add_argument("--repo-root", default=str(repo_root_default))
    parser.add_argument("--excel", default="mig/对外付款申请_20260310214713.xlsx")
    parser.add_argument("--mappings", default="mig/fin_prod_mappings.json")
    parser.add_argument("--manifest", default="mig/attachments/manifest.json")
    parser.add_argument("--manifest-row803", default="mig/attachments/manifest_row803.json")
    parser.add_argument("--att-refs", default="mig/attachments/att_refs.json")
    parser.add_argument("--out-dir", default="mig/sql")
    parser.add_argument("--skip-parity", action="store_true", help="跳过 node amount_cn 一致性烟测")
    args = parser.parse_args()

    repo = Path(args.repo_root).resolve()
    excel_path = repo / args.excel
    out_dir = repo / args.out_dir
    rollback_dir = out_dir / "rollback"
    out_dir.mkdir(parents=True, exist_ok=True)
    rollback_dir.mkdir(parents=True, exist_ok=True)

    mappings = load_json(repo / args.mappings)
    manifest: dict[str, Any] = load_json(repo / args.manifest)
    manifest_row803: dict[str, Any] = load_json(repo / args.manifest_row803)
    att_refs: list[dict[str, Any]] = load_json(repo / args.att_refs)
    resolver = Resolver(mappings)

    voucher_raw_by_id: dict[str, str] = {}
    for ref in att_refs:
        if ref.get("field") == "付款凭证上传" and ref.get("cell"):
            did = ref["data_id"]
            cell = str(ref["cell"]).strip()
            voucher_raw_by_id[did] = (voucher_raw_by_id[did] + "," + cell) if did in voucher_raw_by_id else cell

    rows = load_rows(excel_path)
    by_id: dict[str, dict[str, Any]] = {}
    for row in rows:
        did = as_text(row.get("data_id"))
        if not did:
            raise SystemExit(f"行 {row['__row_no']} 缺 data_id，源文件异常")
        if did in by_id:
            raise SystemExit(f"data_id 重复：{did}")
        by_id[did] = row

    for did in list(EXCLUDED_DATA_IDS) + [DEFERRED_DATA_ID]:
        if did not in by_id:
            raise SystemExit(f"冻结口径引用的 data_id 不在源文件中：{did}")

    included = [r for r in rows if as_text(r.get("data_id")) not in EXCLUDED_DATA_IDS
                and as_text(r.get("data_id")) != DEFERRED_DATA_ID]
    deferred_row = by_id[DEFERRED_DATA_ID]

    # ---- amount_cn 一致性烟测 ----
    all_amounts = [as_amount(r.get("付款金额")) for r in included + [deferred_row]]
    parity: dict[str, Any] = {"skipped": True}
    if not args.skip_parity:
        parity = run_amount_cn_parity(
            repo / "packages/api/src/services/amount-cn.ts",
            [a for a in all_amounts if a is not None],
        )
        if not parity["passed"]:
            raise SystemExit(f"amount_cn 双实现不一致：{json.dumps(parity['mismatches'][:5], ensure_ascii=False)}")

    prevalidation: list[dict[str, Any]] = []
    warnings: list[str] = []
    counts: dict[str, int] = {}

    # ---- 01 payee（归一化键在批内判重，取首次出现原文建档） ----
    payee_first_rows: dict[tuple[str, str], dict[str, Any]] = {}
    for row in included + [deferred_row]:
        name = as_text(row.get("收款单位/收款人"))
        if not name:
            continue  # 进 payment 预校验
        key = (norm_ws(name), norm_ws(as_text(row.get("收款方账号"))))
        payee_first_rows.setdefault(key, row)
        bank = as_text(row.get("收款方开户行信息"))
        if not bank and payee_first_rows[key] is row:
            warnings.append(
                f"data_id={as_text(row.get('data_id'))}（行 {row['__row_no']}）开户行为空，payee.bank_name 落空串，入复核清单"
            )
    payee_stmts = [payee_insert_sql(r, resolver) for r in payee_first_rows.values()]

    # ---- 02 payment ----
    payment_stmts: list[str] = []
    metas: list[dict[str, Any]] = []
    for row in included:
        did = as_text(row.get("data_id"))
        stmt, meta = payment_insert_sql(
            row, resolver, manifest.get(did), voucher_raw_by_id.get(did), prevalidation, warnings
        )
        payment_stmts.append(stmt)
        metas.append(meta)

    status_count = {"paid": 0, "cancelled": 0}
    for m in metas:
        status_count[m["status"]] += 1
    contract_subquery_rows = sum(1 for m in metas if m["contract_subquery"])
    if len(metas) != EXPECT_TOTAL or status_count["paid"] != EXPECT_PAID or status_count["cancelled"] != EXPECT_CANCELLED:
        raise SystemExit(f"行数/状态与冻结口径不符：total={len(metas)} paid={status_count['paid']} cancelled={status_count['cancelled']}")
    if contract_subquery_rows != EXPECT_CONTRACT_LINKED:
        raise SystemExit(f"合同子查询行数 {contract_subquery_rows} != 冻结口径 {EXPECT_CONTRACT_LINKED}")
    if prevalidation:
        warnings.append(f"NOT NULL 预校验命中 {len(prevalidation)} 行——详见 prevalidation_report.md，须处理后再执行")

    # ---- 03 contract attachments ----
    contract_entries = {k: v for k, v in manifest.items() if v.get("domain") == "contract"}
    contract_att_stmts = []
    guard = "(attachments_json IS NULL OR attachments_json='' OR attachments_json='[]')"
    for did in sorted(contract_entries):
        att_json = build_attachments_json(contract_entries[did])
        contract_att_stmts.append(
            f"UPDATE contract SET attachments_json={sql_str(att_json)}, updated_at=updated_at "
            f"WHERE legacy_data_id={sql_str(did)} AND {guard};"
        )

    # ---- 04 paid_amount（固化受影响 contract_no 清单） ----
    affected_contract_nos = sorted({m["contract_no"] for m in metas if m["contract_no"] and m["contract_no"] != ORPHAN_CONTRACT_NO})
    paid_amount_stmts = [
        "UPDATE contract SET paid_amount=(SELECT COALESCE(SUM(amount),0) FROM payment_request "
        f"WHERE contract_id=contract.id AND status='paid'), updated_at=updated_at WHERE contract_no={sql_str(no)};"
        for no in affected_contract_nos
    ]

    # ---- 05 import job ----
    src_md5 = hashlib.md5(excel_path.read_bytes()).hexdigest()[:6].upper()
    today = dt.date.today().strftime("%Y%m%d")
    job_code = f"LGCY-REMOTE-{today}-{src_md5}"
    now = dt.datetime.now(dt.timezone.utc).replace(microsecond=0, tzinfo=None).isoformat() + "Z"
    job_summary = json.dumps(
        {
            "scope": "payment_request 2708 + payee 增量 + contract 附件回填 560 + paid_amount 回填",
            "excluded": list(EXCLUDED_DATA_IDS.values()),
            "deferred": "row803 EX20240801911（90_row803_deferred.sql）",
            "paid": EXPECT_PAID,
            "cancelled": EXPECT_CANCELLED,
            "payee_insert_candidates": len(payee_stmts),
            "contract_attachment_rows": len(contract_att_stmts),
            "paid_amount_contracts": len(affected_contract_nos),
        },
        ensure_ascii=False,
        separators=(",", ":"),
    )
    job_stmt = (
        "INSERT INTO legacy_fin_import_job (job_code, source_file, target_table, status, total_rows, "
        "success_rows, failed_rows, started_at, finished_at, operator_userid, summary_json, created_at, updated_at) "
        f"SELECT {sql_str(job_code)}, {sql_str(excel_path.name)}, 'payment_request', 'completed', {EXPECT_TOTAL}, "
        f"{EXPECT_TOTAL}, 0, {sql_str(now)}, {sql_str(now)}, {sql_str(resolver.fallback_userid)}, {sql_str(job_summary)}, "
        f"{sql_str(now)}, {sql_str(now)} "
        f"WHERE NOT EXISTS (SELECT 1 FROM legacy_fin_import_job WHERE job_code={sql_str(job_code)});"
    )

    # ---- 90 row803 deferred ----
    row803_prevalidation: list[dict[str, Any]] = []
    row803_warnings: list[str] = []
    row803_payee_key = (
        norm_ws(as_text(deferred_row.get("收款单位/收款人"))),
        norm_ws(as_text(deferred_row.get("收款方账号"))),
    )
    row803_stmts: list[str] = []
    if row803_payee_key in payee_first_rows and payee_first_rows[row803_payee_key] is deferred_row:
        # 该收款人仅出现在 row803：payee 建档随 90 文件走，从正批 01 中移除
        payee_stmts = [s for s in payee_stmts if s != payee_insert_sql(deferred_row, resolver)]
        row803_stmts.append(payee_insert_sql(deferred_row, resolver))
    stmt803, meta803 = payment_insert_sql(
        deferred_row, resolver, manifest_row803.get(DEFERRED_DATA_ID),
        voucher_raw_by_id.get(DEFERRED_DATA_ID), row803_prevalidation, row803_warnings,
    )
    row803_stmts.append(stmt803)
    if meta803["contract_no"]:
        row803_stmts.append(
            "UPDATE contract SET paid_amount=(SELECT COALESCE(SUM(amount),0) FROM payment_request "
            "WHERE contract_id=contract.id AND status='paid'), updated_at=updated_at "
            f"WHERE contract_no={sql_str(meta803['contract_no'])};"
        )

    # ---- 写出正向文件 ----
    counts["00_unique_index.sql"] = write_sql(
        out_dir / "00_unique_index.sql",
        ["与 migrations/0082_legacy_fin_unique_index.sql 同内容（IF NOT EXISTS 容错，可重复执行）"],
        [
            "CREATE UNIQUE INDEX IF NOT EXISTS ux_payment_request_legacy_data_id ON payment_request(legacy_data_id) WHERE legacy_data_id IS NOT NULL;",
            "CREATE UNIQUE INDEX IF NOT EXISTS ux_contract_legacy_data_id ON contract(legacy_data_id) WHERE legacy_data_id IS NOT NULL;",
        ],
    )
    write_batched(out_dir, "01_payee", f"payee 增量（legacy_source={PAYEE_LEGACY_SOURCE}；DB 侧归一化谓词判重）", payee_stmts, counts)
    write_batched(out_dir, "02_payment", "payment_request 2708 行（幂等锚点 legacy_data_id；payee/contract 子查询现场解析）", payment_stmts, counts)
    write_batched(out_dir, "03_contract_attachments", "560 条 legacy 合同附件回填（仅空值行）", contract_att_stmts, counts)
    write_batched(out_dir, "04_paid_amount", "paid_amount 回填（受影响 contract_no 固化清单，与 syncContractPaidAmount 同公式）", paid_amount_stmts, counts)
    counts["05_import_job.sql"] = write_sql(out_dir / "05_import_job.sql", [f"导入留档 job_code={job_code}"], [job_stmt])
    counts["90_row803_deferred.sql"] = write_sql(
        out_dir / "90_row803_deferred.sql",
        [
            "row803 EX20240801911 暂缓补导（不入正批）：金额 250000 vs 大写 9909.63 严重不符，",
            "须财务按银行流水确认金额后执行；若确认金额非 250000，需先修改本文件中的 amount/amount_cn。",
            "执行后 CP22020 的 paid_amount 由末条 UPDATE 重算。",
        ],
        row803_stmts,
    )

    # ---- 回滚文件 ----
    rb_counts: dict[str, int] = {}
    rb_counts["00_unique_index.sql"] = write_sql(
        rollback_dir / "00_unique_index.sql",
        ["回滚唯一索引（仅在需要彻底还原时执行；索引本身对既有数据无影响，可保留）"],
        [
            "DROP INDEX IF EXISTS ux_payment_request_legacy_data_id;",
            "DROP INDEX IF EXISTS ux_contract_legacy_data_id;",
        ],
    )
    rb_counts["01_payee.sql"] = write_sql(
        rollback_dir / "01_payee.sql",
        [f"回滚本批新建收款人（锚点 legacy_source={PAYEE_LEGACY_SOURCE}，绝不使用 legacy_data_id 非空——会误伤既有 204 条）"],
        [f"DELETE FROM payee WHERE legacy_source={sql_str(PAYEE_LEGACY_SOURCE)};"],
    )
    rb_payment_stmts = [
        f"DELETE FROM payment_request WHERE legacy_data_id={sql_str(m['data_id'])};" for m in metas
    ]
    write_batched(rollback_dir, "02_payment", "回滚 payment_request（按 legacy_data_id 固化清单）", rb_payment_stmts, rb_counts)
    rb_att_stmts = [
        f"UPDATE contract SET attachments_json=NULL, updated_at=updated_at "
        f"WHERE legacy_data_id={sql_str(did)} AND attachments_json LIKE '%legacy/fin/%';"
        for did in sorted(contract_entries)
    ]
    write_batched(rollback_dir, "03_contract_attachments", "回滚合同附件（导入前实查 560 条原值全 NULL，置 NULL 即恰好还原）", rb_att_stmts, rb_counts)
    rb_paid_stmts = [
        "UPDATE contract SET paid_amount=(SELECT COALESCE(SUM(amount),0) FROM payment_request "
        "WHERE contract_id=contract.id AND status='paid' AND legacy_data_id IS NULL), updated_at=updated_at "
        f"WHERE contract_no={sql_str(no)};"
        for no in affected_contract_nos
    ]
    write_batched(
        rollback_dir, "04_paid_amount",
        "回滚 paid_amount：剔除 legacy 行重算=导入前快照值（多数为 0；CP23109/CP23124 自动还原为新系统已付值）",
        rb_paid_stmts, rb_counts,
    )
    rb_counts["05_import_job.sql"] = write_sql(
        rollback_dir / "05_import_job.sql", ["回滚导入留档"],
        [f"DELETE FROM legacy_fin_import_job WHERE job_code={sql_str(job_code)};"],
    )
    rb_row803 = [
        f"DELETE FROM payment_request WHERE legacy_data_id={sql_str(DEFERRED_DATA_ID)};",
    ]
    if meta803["contract_no"]:
        rb_row803.append(
            "UPDATE contract SET paid_amount=(SELECT COALESCE(SUM(amount),0) FROM payment_request "
            "WHERE contract_id=contract.id AND status='paid'), updated_at=updated_at "
            f"WHERE contract_no={sql_str(meta803['contract_no'])};"
        )
    rb_counts["90_row803_deferred.sql"] = write_sql(
        rollback_dir / "90_row803_deferred.sql",
        ["回滚 row803 补导（payee 若为其独有，由 rollback/01_payee.sql 的 legacy_source 锚点覆盖）"],
        rb_row803,
    )

    # ---- verify.sql ----
    year_paid_sum: dict[str, float] = {}
    for m in metas:
        if m["status"] == "paid" and m["amount"] is not None:
            year_paid_sum[m["year"]] = year_paid_sum.get(m["year"], 0.0) + m["amount"]
    total_paid_amount = round(sum(year_paid_sum.values()), 2)
    verify_stmts = [
        f"SELECT 'A1 付款 legacy 总数(期望 {EXPECT_TOTAL})' AS item, COUNT(*) AS actual FROM payment_request WHERE legacy_data_id IS NOT NULL;",
        "SELECT 'A2 合同 legacy 总数(期望 560，不变)' AS item, COUNT(*) AS actual FROM contract WHERE legacy_data_id IS NOT NULL;",
        f"SELECT 'A3 contract_id 非空(期望 {EXPECT_CONTRACT_LINKED})' AS item, COUNT(*) AS actual FROM payment_request WHERE legacy_data_id IS NOT NULL AND contract_id IS NOT NULL;",
        f"SELECT 'A4 paid(期望 {EXPECT_PAID})' AS item, COUNT(*) AS actual FROM payment_request WHERE legacy_data_id IS NOT NULL AND status='paid';",
        f"SELECT 'A5 cancelled(期望 {EXPECT_CANCELLED})' AS item, COUNT(*) AS actual FROM payment_request WHERE legacy_data_id IS NOT NULL AND status='cancelled';",
        "SELECT 'A6 approved/draft 增量(期望 0)' AS item, COUNT(*) AS actual FROM payment_request WHERE legacy_data_id IS NOT NULL AND status IN ('approved','draft','pending_approval');",
        f"SELECT 'A7 paid 金额合计(期望 {total_paid_amount})' AS item, ROUND(SUM(amount),2) AS actual FROM payment_request WHERE legacy_data_id IS NOT NULL AND status='paid';",
        "SELECT 'A8 分年 paid 金额(对照文件头注释)' AS item, substr(paid_at,1,4) AS year, ROUND(SUM(amount),2) AS actual FROM payment_request WHERE legacy_data_id IS NOT NULL AND status='paid' GROUP BY substr(paid_at,1,4) ORDER BY year;",
        f"SELECT 'B1 本批新建 payee 数(≤{len(payee_stmts)}，命中既有则减)' AS item, COUNT(*) AS actual FROM payee WHERE legacy_source={sql_str(PAYEE_LEGACY_SOURCE)};",
        "SELECT 'B2 payee 解析失配(期望 0)' AS item, COUNT(*) AS actual FROM payment_request WHERE legacy_data_id IS NOT NULL AND payee_id IS NULL;",
        f"SELECT 'B3 contract 子查询失配(期望 0)' AS item, COUNT(*) AS actual FROM payment_request WHERE legacy_data_id IS NOT NULL AND contract_id IS NULL AND legacy_contract_no IS NOT NULL AND legacy_contract_no<>{sql_str(ORPHAN_CONTRACT_NO)};",
        "SELECT 'C1 合同附件回填数(期望 560)' AS item, COUNT(*) AS actual FROM contract WHERE legacy_data_id IS NOT NULL AND attachments_json LIKE '%legacy/fin/%';",
        # 金额正确性由 C3 逐合同 diff 覆盖；此处只验受影响合同全部存在。
        # 不加 paid_amount>0：legacy 付款全为 cancelled 或金额为 0 的合同回填后合法为 0（演练实测 16 个）。
        f"SELECT 'C2 paid_amount 受影响合同数(期望 {len(affected_contract_nos)})' AS item, COUNT(*) AS actual FROM contract WHERE contract_no IN ({', '.join(sql_str(n) for n in affected_contract_nos)});",
        "SELECT 'C3 paid_amount 抽验(逐合同 diff，期望 0 行)' AS item, c.contract_no, c.paid_amount, x.s FROM contract c JOIN (SELECT contract_id, ROUND(COALESCE(SUM(amount),0),2) AS s FROM payment_request WHERE status='paid' GROUP BY contract_id) x ON x.contract_id=c.id WHERE c.legacy_data_id IS NOT NULL AND ROUND(c.paid_amount,2)<>x.s;",
        f"SELECT 'D1 导入留档(期望 1)' AS item, COUNT(*) AS actual FROM legacy_fin_import_job WHERE job_code={sql_str(job_code)};",
    ]
    year_note = "; ".join(f"{y}={round(v, 2)}" for y, v in sorted(year_paid_sum.items()))
    counts["verify.sql"] = write_sql(
        out_dir / "verify.sql",
        [
            "验收查询（方案 §七 1/3/5 及关联项）——期望值内嵌在 item 文案中",
            f"分年 paid 金额期望（按 paid_at 年份，来自 Excel 扣除剔/缓行）：{year_note}",
            "注意：C2/C3 在 row803 补导后需重跑（CP22020 会变化）",
        ],
        verify_stmts,
    )

    # ---- prevalidation_report.md ----
    report_lines = [
        "# NOT NULL 预校验报告（legacy_fin_emit_remote_sql.py 自动生成）",
        "",
        f"- 生成时间：{now}",
        f"- 源文件：{excel_path.name}（md5 前 6：{src_md5}）",
        f"- 本批行数：{len(metas)}（paid {status_count['paid']} / cancelled {status_count['cancelled']}）",
        f"- 剔除：{'; '.join(f'{k}={v}' for k, v in EXCLUDED_DATA_IDS.items())}",
        f"- 暂缓：{DEFERRED_DATA_ID}（row803，90_row803_deferred.sql）",
        "",
        f"## NOT NULL 违规（阻断项）：{len(prevalidation)} 行",
        "",
    ]
    if prevalidation:
        report_lines.append("| 行号 | data_id | 流水号 | 缺失列 |")
        report_lines.append("|---|---|---|---|")
        for item in prevalidation:
            report_lines.append(f"| {item['row_no']} | {item['data_id']} | {item['serial_no'] or ''} | {', '.join(item['missing'])} |")
    else:
        report_lines.append("无。3 个问题行已按冻结口径剔/缓，其余行 NOT NULL 目标列全部有值。")
    report_lines += ["", f"## 提示（非阻断，复核清单）：{len(warnings)} 条", ""]
    report_lines += [f"- {w}" for w in warnings] or ["- 无"]
    if row803_prevalidation:
        report_lines += ["", "## row803（暂缓文件）预校验", ""]
        report_lines += [f"- {json.dumps(i, ensure_ascii=False)}" for i in row803_prevalidation]
    (out_dir / "prevalidation_report.md").write_text("\n".join(report_lines) + "\n", encoding="utf-8")

    # ---- stats ----
    stats = {
        "generated_at": now,
        "job_code": job_code,
        "files": {**counts, "rollback": rb_counts},
        "payment_rows": len(metas),
        "status_count": status_count,
        "payee_insert_candidates": len(payee_stmts),
        "payee_subquery_rows": len(metas) + 1,
        "contract_subquery_rows": contract_subquery_rows,
        "paid_amount_contracts": len(affected_contract_nos),
        "contract_attachment_rows": len(contract_att_stmts),
        "payment_attachment_rows": sum(1 for m in metas if manifest.get(m["data_id"])),
        "resolver": resolver.stats,
        "amount_cn_parity": parity,
        "prevalidation_errors": len(prevalidation),
        "prevalidation_warnings": len(warnings),
        "year_paid_sum": {y: round(v, 2) for y, v in sorted(year_paid_sum.items())},
        "total_paid_amount": total_paid_amount,
    }
    (out_dir / "emit_stats.json").write_text(json.dumps(stats, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(stats, ensure_ascii=False, indent=2))
    if prevalidation:
        raise SystemExit("NOT NULL 预校验未通过，见 prevalidation_report.md")


if __name__ == "__main__":
    main()
