#!/usr/bin/env python3
from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import json
import sqlite3
import uuid
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def now_iso() -> str:
    return dt.datetime.utcnow().replace(microsecond=0).isoformat() + "Z"


def as_text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, str):
        text = value.strip()
        return text if text else None
    text = str(value).strip()
    return text if text else None


def as_amount(value: Any) -> float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = as_text(value)
    if not text:
        return None
    cleaned = text.replace(",", "").replace("，", "")
    try:
        return float(cleaned)
    except ValueError:
        return None


def as_iso_datetime(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, dt.datetime):
        return value.replace(microsecond=0).isoformat() + "Z"
    if isinstance(value, dt.date):
        return dt.datetime(value.year, value.month, value.day).isoformat() + "Z"
    text = as_text(value)
    if not text:
        return None
    normalized = text.replace("/", "-")
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d"):
        try:
            parsed = dt.datetime.strptime(normalized, fmt)
            return parsed.isoformat() + "Z"
        except ValueError:
            continue
    return text


def split_file_keys(value: Any) -> list[str]:
    text = as_text(value)
    if not text:
        return []
    items = [item.strip() for item in text.split(",")]
    return [item for item in items if item]


def normalize_bank_account(value: str | None) -> str | None:
    if not value:
        return None
    text = value.replace(" ", "").replace("\t", "")
    if text.endswith(".0") and text[:-2].isdigit():
        text = text[:-2]
    return text or None


def stable_legacy_code(prefix: str, seed: str) -> str:
    digest = hashlib.sha1(seed.encode("utf-8")).hexdigest()[:16].upper()
    return f"{prefix}-{digest}"


def map_contract_status(value: Any) -> str:
    text = as_text(value)
    if text == "流转完成":
        return "active"
    if text == "进行中":
        return "pending_approval"
    return "draft"


def map_payment_status(value: Any, paid_at: str | None) -> str:
    text = as_text(value)
    if text == "流转完成":
        return "paid" if paid_at else "approved"
    if text == "进行中":
        return "pending_approval"
    return "draft"


def map_reimbursement_envelope_status(value: Any) -> str:
    text = as_text(value)
    if text == "已完成":
        return "approved"
    if text == "集中申报":
        return "draft"
    return "draft"


def load_rows(excel_path: Path) -> list[dict[str, Any]]:
    wb = load_workbook(excel_path, data_only=True)
    ws = wb.active
    header_row = 1
    best_count = 0
    for idx in range(1, min(ws.max_row, 20) + 1):
        values = [as_text(ws.cell(idx, col_idx).value) for col_idx in range(1, ws.max_column + 1)]
        count = sum(1 for value in values if value)
        if count > best_count:
            best_count = count
            header_row = idx

    header = [as_text(ws.cell(header_row, col_idx).value) or "" for col_idx in range(1, ws.max_column + 1)]
    rows: list[dict[str, Any]] = []
    for row_idx in range(header_row + 1, ws.max_row + 1):
        values = [ws.cell(row_idx, col_idx).value for col_idx in range(1, len(header) + 1)]
        if not any(v is not None and as_text(v) is not None for v in values):
            continue
        rows.append({"__row_no": row_idx, **{header[i]: values[i] for i in range(len(header)) if header[i]}})
    return rows


@dataclass
class JobStats:
    total: int = 0
    success: int = 0
    failed: int = 0


class Migrator:
    def __init__(self, db: sqlite3.Connection | None, mapping: dict[str, Any], mode: str):
        self.db = db
        self.mapping = mapping
        self.mode = mode
        self.summary: dict[str, Any] = {"mode": mode, "generated_at": now_iso(), "tables": {}, "errors": []}
        self.job_ids: dict[str, int] = {}
        self.max_error_samples = 200
        self.dropped_error_samples = 0

    def default_userid(self) -> str | None:
        value = self.mapping.get("default_userid")
        if value is None:
            return None
        text = str(value).strip()
        return text if text else None

    def default_project_id(self) -> int | None:
        value = self.mapping.get("default_project_id")
        if value is None:
            return None
        return int(value)

    def error(self, table: str, row_no: int | None, legacy_data_id: str | None, code: str, message: str, raw: dict[str, Any]) -> None:
        payload = {
            "table": table,
            "row_no": row_no,
            "legacy_data_id": legacy_data_id,
            "error_code": code,
            "error_message": message,
        }
        if len(self.summary["errors"]) < self.max_error_samples:
            self.summary["errors"].append(payload)
        else:
            self.dropped_error_samples += 1
        if self.mode != "apply" or self.db is None:
            return
        job_id = self.job_ids.get(table)
        if not job_id:
            return
        self.db.execute(
            """
            INSERT INTO legacy_fin_import_error
              (job_id, row_no, legacy_data_id, error_code, error_message, raw_row_json, created_at)
            VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            (
                job_id,
                row_no,
                legacy_data_id,
                code,
                message,
                json.dumps(raw, ensure_ascii=False, default=str),
                now_iso(),
            ),
        )

    def start_job(self, table: str, source_file: str) -> None:
        if self.mode != "apply" or self.db is None:
            return
        stamp = dt.datetime.utcnow().strftime("%Y%m%d%H%M%S%f")
        source_tag = hashlib.md5(source_file.encode("utf-8")).hexdigest()[:6].upper()
        job_code = f"LGCY-{table}-{source_tag}-{stamp}"
        now = now_iso()
        cur = self.db.execute(
            """
            INSERT INTO legacy_fin_import_job
              (job_code, source_file, target_table, status, started_at, created_at, updated_at)
            VALUES (?, ?, ?, 'running', ?, ?, ?)
            """,
            (job_code, source_file, table, now, now, now),
        )
        self.job_ids[table] = int(cur.lastrowid)

    def finish_job(self, table: str, stats: JobStats) -> None:
        existing = self.summary["tables"].get(table, {"total": 0, "success": 0, "failed": 0})
        self.summary["tables"][table] = {
            "total": int(existing["total"]) + stats.total,
            "success": int(existing["success"]) + stats.success,
            "failed": int(existing["failed"]) + stats.failed,
        }
        if self.mode != "apply" or self.db is None:
            return
        job_id = self.job_ids.get(table)
        if not job_id:
            return
        status = "completed" if stats.failed == 0 else ("partial_failed" if stats.success > 0 else "failed")
        self.db.execute(
            """
            UPDATE legacy_fin_import_job
            SET status = ?,
                total_rows = ?,
                success_rows = ?,
                failed_rows = ?,
                finished_at = ?,
                summary_json = ?,
                updated_at = ?
            WHERE id = ?
            """,
            (
                status,
                stats.total,
                stats.success,
                stats.failed,
                now_iso(),
                json.dumps(self.summary["tables"][table], ensure_ascii=False),
                now_iso(),
                job_id,
            ),
        )

    def resolve_project_id(self, legacy_name: str | None, allow_default: bool = False) -> int | None:
        if not legacy_name:
            return self.default_project_id() if allow_default else None
        project_map = self.mapping.get("project_map", {})
        if legacy_name in project_map:
            mapped = project_map[legacy_name]
            if mapped is None:
                return self.default_project_id() if allow_default else None
            return int(mapped)
        if self.db is None:
            return self.default_project_id() if allow_default else None
        row = self.db.execute(
            """
            SELECT id FROM project
            WHERE project_code = ? OR project_name = ?
            ORDER BY id DESC LIMIT 1
            """,
            (legacy_name, legacy_name),
        ).fetchone()
        if row:
            return int(row[0])
        return self.default_project_id() if allow_default else None

    def resolve_userid(self, user_name: str | None, allow_default: bool = False) -> str | None:
        if not user_name:
            return self.default_userid() if allow_default else None
        user_map = self.mapping.get("user_map", {})
        if user_name in user_map:
            mapped = user_map[user_name]
            if mapped is None:
                return self.default_userid() if allow_default else None
            return str(mapped)
        if allow_default:
            return self.default_userid() or user_name
        return user_name

    def resolve_entity_id(self, entity_name: str | None) -> int | None:
        if not entity_name:
            return None
        entity_map = self.mapping.get("entity_map", {})
        if entity_name in entity_map:
            mapped = entity_map[entity_name]
            if mapped is None:
                return None
            return int(mapped)
        if self.db is None:
            return None
        row = self.db.execute(
            "SELECT id FROM entity WHERE entity_name = ? ORDER BY id DESC LIMIT 1",
            (entity_name,),
        ).fetchone()
        return int(row[0]) if row else None

    def should_auto_create_entity(self) -> bool:
        return bool(self.mapping.get("auto_create_entity", False))

    def ensure_entity_id(self, entity_name: str | None) -> int | None:
        entity_id = self.resolve_entity_id(entity_name)
        if entity_id is not None:
            return entity_id
        if not self.should_auto_create_entity() or not entity_name or self.db is None:
            return None
        now = now_iso()
        entity_code = stable_legacy_code("AUTENT", entity_name)
        try:
            self.db.execute(
                """
                INSERT INTO entity
                  (entity_code, entity_name, legal_person, status, created_at, updated_at,
                   legacy_data_id, tax_id, legacy_source, legacy_submitted_by, legacy_submitted_at)
                VALUES (?, ?, NULL, 'active', ?, ?, NULL, NULL, 'fin_data_auto', ?, ?)
                """,
                (
                    entity_code,
                    entity_name,
                    now,
                    now,
                    self.default_userid() or "legacy_import",
                    now,
                ),
            )
        except sqlite3.IntegrityError:
            existing = self.db.execute("SELECT id FROM entity WHERE entity_name = ? LIMIT 1", (entity_name,)).fetchone()
            if existing:
                return int(existing[0])
            alt_code = stable_legacy_code("AUTENT", f"{entity_name}|fallback")
            self.db.execute(
                """
                INSERT INTO entity
                  (entity_code, entity_name, legal_person, status, created_at, updated_at,
                   legacy_data_id, tax_id, legacy_source, legacy_submitted_by, legacy_submitted_at)
                VALUES (?, ?, NULL, 'active', ?, ?, NULL, NULL, 'fin_data_auto', ?, ?)
                """,
                (
                    alt_code,
                    entity_name,
                    now,
                    now,
                    self.default_userid() or "legacy_import",
                    now,
                ),
            )
        row = self.db.execute("SELECT id FROM entity WHERE entity_code = ? LIMIT 1", (entity_code,)).fetchone()
        return int(row[0]) if row else None

    def upsert_payee(self, row: dict[str, Any], payee_type: str) -> bool:
        legacy_data_id = as_text(row.get("data_id"))
        name = as_text(row.get("收款单位") or row.get("收款人姓名"))
        bank_name = as_text(row.get("收款方开户行"))
        bank_account = as_text(row.get("收款方账户"))
        phone = as_text(row.get("电话"))
        submitter = self.resolve_userid(as_text(row.get("提交人")), allow_default=True) or "legacy_import"
        submitted_at = as_iso_datetime(row.get("提交时间")) or now_iso()
        updated_at = as_iso_datetime(row.get("更新时间")) or submitted_at

        if not name or not bank_name or not bank_account:
            return False

        if self.mode == "apply" and self.db is not None:
            existed = None
            if legacy_data_id:
                existed = self.db.execute(
                    "SELECT id FROM payee WHERE legacy_data_id = ? LIMIT 1",
                    (legacy_data_id,),
                ).fetchone()
            if not existed:
                existed = self.db.execute(
                    "SELECT id FROM payee WHERE name = ? AND bank_account = ? LIMIT 1",
                    (name, bank_account),
                ).fetchone()
            if existed:
                return True

            self.db.execute(
                """
                INSERT INTO payee
                  (name, bank_name, bank_account, phone, type, status, created_by_userid, created_at, updated_at,
                   legacy_data_id, legacy_source, legacy_submitted_by, legacy_submitted_at)
                VALUES (?, ?, ?, ?, ?, 'active', ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    name,
                    bank_name,
                    bank_account,
                    phone,
                    payee_type,
                    submitter,
                    submitted_at,
                    updated_at,
                    legacy_data_id,
                    "fin_data",
                    as_text(row.get("提交人")),
                    submitted_at,
                ),
            )
        return True

    def upsert_entity(self, row: dict[str, Any]) -> bool:
        legacy_data_id = as_text(row.get("data_id"))
        entity_name = as_text(row.get("公司名称"))
        tax_id = as_text(row.get("税号"))
        submitter = self.resolve_userid(as_text(row.get("提交人")), allow_default=True) or "legacy_import"
        created_at = as_iso_datetime(row.get("提交时间")) or now_iso()
        updated_at = as_iso_datetime(row.get("更新时间")) or created_at
        if not entity_name:
            return False

        entity_code = stable_legacy_code("LEGENT", f"{entity_name}|{tax_id or ''}")

        if self.mode == "apply" and self.db is not None:
            existed = None
            if legacy_data_id:
                existed = self.db.execute(
                    "SELECT id FROM entity WHERE legacy_data_id = ? LIMIT 1",
                    (legacy_data_id,),
                ).fetchone()
            if not existed and tax_id:
                existed = self.db.execute("SELECT id FROM entity WHERE tax_id = ? LIMIT 1", (tax_id,)).fetchone()
            if not existed:
                existed = self.db.execute("SELECT id FROM entity WHERE entity_name = ? LIMIT 1", (entity_name,)).fetchone()
            if existed:
                return True

            self.db.execute(
                """
                INSERT INTO entity
                  (entity_code, entity_name, legal_person, status, created_at, updated_at,
                   legacy_data_id, tax_id, legacy_source, legacy_submitted_by, legacy_submitted_at)
                VALUES (?, ?, NULL, 'active', ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    entity_code,
                    entity_name,
                    created_at,
                    updated_at,
                    legacy_data_id,
                    tax_id,
                    "fin_data",
                    submitter,
                    created_at,
                ),
            )
        return True

    def process_payee(self, rows: list[dict[str, Any]], payee_type: str, source_file: str) -> None:
        table = "payee"
        self.start_job(table, source_file)
        stats = JobStats()
        for row in rows:
            stats.total += 1
            if self.upsert_payee(row, payee_type):
                stats.success += 1
            else:
                stats.failed += 1
                self.error(table, row.get("__row_no"), as_text(row.get("data_id")), "PAYEE_REQUIRED_MISSING", "missing payee name/bank/account", row)
        self.finish_job(table, stats)

    def process_entity(self, rows: list[dict[str, Any]], source_file: str) -> None:
        table = "entity"
        self.start_job(table, source_file)
        stats = JobStats()
        for row in rows:
            stats.total += 1
            if self.upsert_entity(row):
                stats.success += 1
            else:
                stats.failed += 1
                self.error(table, row.get("__row_no"), as_text(row.get("data_id")), "ENTITY_REQUIRED_MISSING", "missing entity_name", row)
        self.finish_job(table, stats)

    def process_contract(self, rows: list[dict[str, Any]], source_file: str) -> None:
        table = "contract"
        self.start_job(table, source_file)
        stats = JobStats()
        for row in rows:
            stats.total += 1
            legacy_data_id = as_text(row.get("data_id"))
            contract_no = as_text(row.get("合同号"))
            project_name = as_text(row.get("所属工程"))
            full_project_name = as_text(row.get("工程全称"))
            project_id = self.resolve_project_id(project_name, allow_default=True) or self.resolve_project_id(
                full_project_name, allow_default=True
            )
            entity_id = self.ensure_entity_id(as_text(row.get("合同乙方")))
            handler = self.resolve_userid(as_text(row.get("经办人")), allow_default=True)
            contract_date = as_iso_datetime(row.get("合同日期")) or now_iso()
            amount = as_amount(row.get("合同金额"))
            missing: list[str] = []
            if not contract_no:
                missing.append("contract_no")
            if project_id is None:
                missing.append("project_id")
            if entity_id is None:
                missing.append("entity_id")
            if not handler:
                missing.append("handler_userid")
            if amount is None:
                missing.append("total_amount")
            if missing:
                stats.failed += 1
                self.error(
                    table,
                    row.get("__row_no"),
                    legacy_data_id,
                    "CONTRACT_MAPPING_FAILED",
                    f"missing {','.join(missing)}",
                    row,
                )
                continue

            if self.mode == "apply" and self.db is not None:
                existed = None
                if legacy_data_id:
                    existed = self.db.execute("SELECT id FROM contract WHERE legacy_data_id = ? LIMIT 1", (legacy_data_id,)).fetchone()
                if not existed:
                    existed = self.db.execute("SELECT id FROM contract WHERE contract_no = ? LIMIT 1", (contract_no,)).fetchone()
                if not existed:
                    self.db.execute(
                        """
                        INSERT INTO contract
                          (contract_no, direction, project_id, entity_id, counterparty, contract_type,
                           handler_userid, contract_date, total_amount, paid_amount, invoiced_amount,
                           received_amount, status, ai_review_status, attachments_json, remark,
                           created_at, updated_at, legacy_data_id, legacy_project_name,
                           legacy_project_full_name, legacy_contract_status, legacy_current_node,
                           legacy_current_owner, legacy_submitter_name, legacy_submitted_at,
                           legacy_contract_party_a, legacy_contract_party_b)
                        VALUES (?, 'outbound', ?, ?, ?, ?, ?, ?, ?, 0, 0, 0, ?, 'none', ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                        """,
                        (
                            contract_no,
                            project_id,
                            entity_id,
                            as_text(row.get("合同甲方")) or "legacy_counterparty",
                            as_text(row.get("合同类型")) or "其他",
                            handler,
                            contract_date,
                            amount,
                            map_contract_status(row.get("流程状态")),
                            json.dumps(split_file_keys(row.get("合同扫描件")), ensure_ascii=False),
                            as_text(row.get("合同内容")),
                            as_iso_datetime(row.get("提交时间")) or now_iso(),
                            as_iso_datetime(row.get("更新时间")) or now_iso(),
                            legacy_data_id,
                            project_name,
                            full_project_name,
                            as_text(row.get("流程状态")),
                            as_text(row.get("当前节点")),
                            as_text(row.get("当前负责人")),
                            as_text(row.get("提交人")),
                            as_iso_datetime(row.get("提交时间")),
                            as_text(row.get("合同甲方")),
                            as_text(row.get("合同乙方")),
                        ),
                    )
            stats.success += 1
        self.finish_job(table, stats)

    def resolve_payee_id(self, name: str | None, account: str | None) -> int | None:
        if self.db is None or not name:
            return None
        normalized_input = normalize_bank_account(account)

        rows = self.db.execute(
            "SELECT id, bank_account FROM payee WHERE name = ? AND status = 'active' ORDER BY id DESC",
            (name,),
        ).fetchall()
        if not rows:
            return None

        if normalized_input:
            for row in rows:
                if normalize_bank_account(as_text(row[1])) == normalized_input:
                    return int(row[0])

        if len(rows) == 1:
            return int(rows[0][0])

        return None

    def resolve_contract_id(self, contract_no: str | None) -> int | None:
        if self.db is None or not contract_no:
            return None
        row = self.db.execute("SELECT id FROM contract WHERE contract_no = ? ORDER BY id DESC LIMIT 1", (contract_no,)).fetchone()
        return int(row[0]) if row else None

    def should_auto_create_payee(self) -> bool:
        return bool(self.mapping.get("auto_create_payee", True))

    def ensure_payee_id(
        self,
        name: str | None,
        bank_name: str | None,
        bank_account: str | None,
        phone: str | None,
        payee_type: str = "external",
    ) -> int | None:
        payee_id = self.resolve_payee_id(name, bank_account)
        if payee_id is not None:
            return payee_id
        if self.db is None or not self.should_auto_create_payee() or not name or not bank_name or not bank_account:
            return None

        now = now_iso()
        creator = self.default_userid() or "legacy_import"
        self.db.execute(
            """
            INSERT INTO payee
              (name, bank_name, bank_account, phone, type, status, created_by_userid, created_at, updated_at,
               legacy_data_id, legacy_source, legacy_submitted_by, legacy_submitted_at)
            VALUES (?, ?, ?, ?, ?, 'active', ?, ?, ?, NULL, 'fin_data_auto', ?, ?)
            """,
            (name, bank_name, bank_account, phone, payee_type, creator, now, now, creator, now),
        )
        row = self.db.execute(
            "SELECT id FROM payee WHERE name = ? AND bank_account = ? ORDER BY id DESC LIMIT 1",
            (name, bank_account),
        ).fetchone()
        return int(row[0]) if row else None

    def process_payment_request(self, rows: list[dict[str, Any]], source_file: str) -> None:
        table = "payment_request"
        self.start_job(table, source_file)
        stats = JobStats()
        for row in rows:
            stats.total += 1
            legacy_data_id = as_text(row.get("data_id"))
            serial_no = as_text(row.get("流水号"))
            project_name = as_text(row.get("所属项目/工程"))
            full_project_name = as_text(row.get("工程全称"))
            project_id = self.resolve_project_id(project_name, allow_default=True) or self.resolve_project_id(
                full_project_name, allow_default=True
            )
            applicant_name = as_text(row.get("申请人"))
            applicant_userid = self.resolve_userid(applicant_name, allow_default=True)
            payee_name = as_text(row.get("收款单位/收款人"))
            payee_bank_name = as_text(row.get("收款方开户行信息"))
            payee_account = as_text(row.get("收款方账号"))
            payee_phone = as_text(row.get("收款方电话"))
            payee_id = self.ensure_payee_id(payee_name, payee_bank_name, payee_account, payee_phone, payee_type="external")
            contract_no = as_text(row.get("合同号"))
            contract_id = self.resolve_contract_id(contract_no)
            amount = as_amount(row.get("付款金额"))
            expected_payment_date = as_iso_datetime(row.get("预计付款日期")) or now_iso()
            paid_at = as_iso_datetime(row.get("实际付款日期"))

            missing: list[str] = []
            if not serial_no:
                missing.append("serial_no")
            if project_id is None:
                missing.append("project_id")
            if not applicant_userid:
                missing.append("applicant_userid")
            if payee_id is None:
                missing.append("payee_id")
            if amount is None:
                missing.append("amount")
            if missing:
                stats.failed += 1
                self.error(
                    table,
                    row.get("__row_no"),
                    legacy_data_id,
                    "PAYMENT_MAPPING_FAILED",
                    f"missing {','.join(missing)}",
                    row,
                )
                continue

            if self.mode == "apply" and self.db is not None:
                existed = None
                if legacy_data_id:
                    existed = self.db.execute(
                        "SELECT id FROM payment_request WHERE legacy_data_id = ? LIMIT 1",
                        (legacy_data_id,),
                    ).fetchone()
                if not existed:
                    existed = self.db.execute("SELECT id FROM payment_request WHERE serial_no = ? LIMIT 1", (serial_no,)).fetchone()
                if not existed:
                    self.db.execute(
                        """
                        INSERT INTO payment_request
                          (serial_no, contract_id, project_id, applicant_userid, expected_payment_date,
                           category, payee_id, amount, amount_cn, content, payment_method, status,
                           approval_instance_id, ai_risk_status, ai_risk_result, attachments_json,
                           remark, created_at, updated_at, paid_at,
                           legacy_data_id, legacy_applicant_name, legacy_project_name,
                           legacy_project_full_name, legacy_contract_no, legacy_contract_status,
                           legacy_invoice_type, legacy_payee_name, legacy_payee_bank_name,
                           legacy_payee_account, legacy_payee_phone, legacy_handler_name,
                           legacy_specified_pay_flag, legacy_specified_pay_date,
                           legacy_voucher_keys_raw, legacy_flow_status, legacy_current_node,
                           legacy_current_owner, legacy_submitter_name, legacy_submitted_at)
                        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, NULL, 'none', NULL, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                        """,
                        (
                            serial_no,
                            contract_id,
                            project_id,
                            applicant_userid,
                            expected_payment_date,
                            as_text(row.get("报销类别")) or "其他",
                            payee_id,
                            amount,
                            as_text(row.get("付款金额大写")) or "",
                            as_text(row.get("付款内容（详情）")) or "legacy_import",
                            as_text(row.get("付款方式")) or "对公转账",
                            map_payment_status(row.get("流程状态"), paid_at),
                            json.dumps(split_file_keys(row.get("上传发票")), ensure_ascii=False),
                            as_text(row.get("备注")),
                            as_iso_datetime(row.get("提交时间")) or now_iso(),
                            as_iso_datetime(row.get("更新时间")) or now_iso(),
                            paid_at,
                            legacy_data_id,
                            applicant_name,
                            project_name,
                            full_project_name,
                            contract_no,
                            as_text(row.get("合同状态")),
                            as_text(row.get("发票类型")),
                            payee_name,
                            as_text(row.get("收款方开户行信息")),
                            payee_account,
                            as_text(row.get("收款方电话")),
                            as_text(row.get("办理人")),
                            as_text(row.get("指定日期付款")),
                            as_iso_datetime(row.get("指定付款日期")),
                            as_text(row.get("付款凭证上传")),
                            as_text(row.get("流程状态")),
                            as_text(row.get("当前节点")),
                            as_text(row.get("当前负责人")),
                            as_text(row.get("提交人")),
                            as_iso_datetime(row.get("提交时间")),
                        ),
                    )
            stats.success += 1
        self.finish_job(table, stats)

    def process_reimbursement_envelope(self, rows: list[dict[str, Any]], source_file: str) -> None:
        table = "reimbursement_envelope"
        self.start_job(table, source_file)
        stats = JobStats()
        for row in rows:
            stats.total += 1
            legacy_data_id = as_text(row.get("data_id"))
            envelope_no = as_text(row.get("流水号")) or stable_legacy_code("ENV", legacy_data_id or str(row.get("__row_no")))
            project_name = as_text(row.get("所属工程"))
            project_id = self.resolve_project_id(project_name, allow_default=True)
            applicant_name = as_text(row.get("提交人"))
            applicant_userid = self.resolve_userid(applicant_name, allow_default=True)
            amount = as_amount(row.get("报销金额"))
            missing: list[str] = []
            if project_id is None:
                missing.append("project_id")
            if not applicant_userid:
                missing.append("applicant_userid")
            if amount is None:
                missing.append("amount")
            if missing:
                stats.failed += 1
                self.error(
                    table,
                    row.get("__row_no"),
                    legacy_data_id,
                    "ENVELOPE_MAPPING_FAILED",
                    f"missing {','.join(missing)}",
                    row,
                )
                continue

            envelope_uuid = str(uuid.uuid5(uuid.NAMESPACE_DNS, f"legacy-envelope:{legacy_data_id or envelope_no}"))
            attachments = split_file_keys(row.get("附件上传"))

            if self.mode == "apply" and self.db is not None:
                existed = None
                if legacy_data_id:
                    existed = self.db.execute(
                        "SELECT id FROM reimbursement_envelope WHERE legacy_data_id = ? LIMIT 1",
                        (legacy_data_id,),
                    ).fetchone()
                if not existed:
                    existed = self.db.execute(
                        "SELECT id FROM reimbursement_envelope WHERE envelope_no = ? LIMIT 1",
                        (envelope_no,),
                    ).fetchone()
                if existed:
                    envelope_id = int(existed[0])
                else:
                    cur = self.db.execute(
                        """
                        INSERT INTO reimbursement_envelope
                          (envelope_no, envelope_uuid, project_id, applicant_userid, channel, invoice_no,
                           expense_date, amount, category, description, status, precheck_status,
                           precheck_comment, verification_status, verification_note, batch_id,
                           approval_instance_id, pdf_task_id, pdf_status, pdf_file_key, ocr_task_id,
                           ocr_status, duplicate_check_status, submitted_at, received_at,
                           created_at, updated_at, legacy_data_id, legacy_serial_no, legacy_project_name,
                           legacy_ticket_type, legacy_invoice_received_flag, legacy_ticket_match_flag,
                           legacy_reimbursement_status, legacy_flow_status, legacy_current_node,
                           legacy_current_owner, legacy_submitter_name, legacy_submitted_at)
                        VALUES (?, ?, ?, ?, 'physical', NULL, ?, ?, ?, ?, ?, 'none', NULL, 'pending', NULL, NULL,
                                NULL, NULL, 'none', NULL, NULL, 'none', 'none', ?, NULL,
                                ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                        """,
                        (
                            envelope_no,
                            envelope_uuid,
                            project_id,
                            applicant_userid,
                            as_iso_datetime(row.get("费用发生日期")),
                            amount,
                            as_text(row.get("报销类别")),
                            as_text(row.get("报销内容")),
                            map_reimbursement_envelope_status(row.get("流程状态")),
                            as_iso_datetime(row.get("提交时间")),
                            as_iso_datetime(row.get("提交时间")) or now_iso(),
                            as_iso_datetime(row.get("更新时间")) or now_iso(),
                            legacy_data_id,
                            as_text(row.get("流水号")),
                            project_name,
                            as_text(row.get("票据类型")),
                            as_text(row.get("是否收到发票")),
                            as_text(row.get("票据信息与报销信息是否一致")),
                            as_text(row.get("报销状态")),
                            as_text(row.get("流程状态")),
                            as_text(row.get("当前节点")),
                            as_text(row.get("当前负责人")),
                            applicant_name,
                            as_iso_datetime(row.get("提交时间")),
                        ),
                    )
                    envelope_id = int(cur.lastrowid)

                for key in attachments:
                    self.db.execute(
                        """
                        INSERT OR IGNORE INTO envelope_attachment
                          (envelope_id, attachment_type, file_key, file_name, file_size, invoice_no, amount, issue_date, status, created_at)
                        VALUES (?, 'invoice', ?, NULL, NULL, NULL, NULL, NULL, 'active', ?)
                        """,
                        (envelope_id, key, now_iso()),
                    )

            stats.success += 1
        self.finish_job(table, stats)

    def process_reimbursement_batch(self, rows: list[dict[str, Any]], source_file: str) -> None:
        table = "reimbursement_batch"
        self.start_job(table, source_file)
        stats = JobStats()
        envelope_by_serial: dict[str, tuple[int, float]] = {}
        if self.mode == "apply" and self.db is not None:
            env_rows = self.db.execute(
                """
                SELECT id, IFNULL(amount, 0) as amount, legacy_serial_no, envelope_no
                FROM reimbursement_envelope
                WHERE status != 'deleted'
                ORDER BY id DESC
                """
            ).fetchall()
            for env_row in env_rows:
                env_id = int(env_row[0])
                env_amount = float(env_row[1] or 0)
                legacy_serial = as_text(env_row[2])
                envelope_no = as_text(env_row[3])
                if legacy_serial and legacy_serial not in envelope_by_serial:
                    envelope_by_serial[legacy_serial] = (env_id, env_amount)
                if envelope_no and envelope_no not in envelope_by_serial:
                    envelope_by_serial[envelope_no] = (env_id, env_amount)

        current_parent: dict[str, Any] | None = None

        def flush_parent(parent: dict[str, Any]) -> None:
            parent_row: dict[str, Any] = parent["row"]
            child_rows: list[dict[str, Any]] = parent["child_rows"]
            details: list[str] = parent["details"]

            legacy_data_id = as_text(parent_row.get("data_id"))
            project_name = as_text(parent_row.get("所属项目/工程"))
            project_id = self.resolve_project_id(project_name, allow_default=True)
            claimant_name = as_text(parent_row.get("待报销人"))
            applicant_userid = self.resolve_userid(claimant_name, allow_default=True)
            total_amount = as_amount(parent_row.get("报销总金额"))

            missing: list[str] = []
            if project_id is None:
                missing.append("project_id")
            if not applicant_userid:
                missing.append("applicant_userid")
            if total_amount is None:
                missing.append("total_amount")
            if not legacy_data_id:
                missing.append("legacy_data_id")

            if missing:
                stats.failed += 1 + len(child_rows)
                self.error(
                    table,
                    parent_row.get("__row_no"),
                    legacy_data_id,
                    "BATCH_MAPPING_FAILED",
                    f"missing {','.join(missing)}",
                    parent_row,
                )
                return

            matched_envelopes: list[tuple[int, float, str]] = []
            unmatched_serials: list[str] = []
            seen_serials: set[str] = set()
            for detail in details:
                if detail in seen_serials:
                    continue
                seen_serials.add(detail)
                matched = envelope_by_serial.get(detail)
                if matched is None:
                    unmatched_serials.append(detail)
                    continue
                matched_envelopes.append((matched[0], matched[1], detail))

            if self.mode == "apply" and self.db is not None:
                batch_no = stable_legacy_code("LGB", legacy_data_id)
                existed = self.db.execute(
                    "SELECT id FROM reimbursement_batch WHERE legacy_data_id = ? LIMIT 1",
                    (legacy_data_id,),
                ).fetchone()

                details_payload = {
                    "primary_detail": as_text(parent_row.get("报销详情")),
                    "all_details": details,
                    "matched_details": [item[2] for item in matched_envelopes],
                    "unmatched_details": unmatched_serials,
                }
                details_text = json.dumps(details_payload, ensure_ascii=False, default=str)

                if not existed:
                    cur = self.db.execute(
                        """
                        INSERT INTO reimbursement_batch
                          (batch_no, project_id, applicant_userid, total_amount, item_count, status,
                           ai_audit_status, ai_audit_result, approval_instance_id, created_at, updated_at,
                           legacy_data_id, legacy_claimant_name, legacy_project_name, legacy_details_raw,
                           legacy_amount_cn, legacy_payment_method, legacy_handler_name,
                           legacy_specified_pay_flag, legacy_specified_pay_date, legacy_voucher_keys_raw,
                           legacy_flow_status, legacy_current_node, legacy_current_owner,
                           legacy_submitter_name, legacy_submitted_at)
                        VALUES (?, ?, ?, ?, 0, ?, 'none', NULL, NULL, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                        """,
                        (
                            batch_no,
                            project_id,
                            applicant_userid,
                            total_amount,
                            map_reimbursement_envelope_status(parent_row.get("流程状态")),
                            as_iso_datetime(parent_row.get("提交时间")) or now_iso(),
                            as_iso_datetime(parent_row.get("更新时间")) or now_iso(),
                            legacy_data_id,
                            claimant_name,
                            project_name,
                            details_text,
                            as_text(parent_row.get("总金额大写")),
                            as_text(parent_row.get("付款方式")),
                            as_text(parent_row.get("办理人")),
                            as_text(parent_row.get("指定日期付款")),
                            as_iso_datetime(parent_row.get("指定付款日期")),
                            as_text(parent_row.get("付款凭证上传")),
                            as_text(parent_row.get("流程状态")),
                            as_text(parent_row.get("当前节点")),
                            as_text(parent_row.get("当前负责人")),
                            as_text(parent_row.get("提交人")),
                            as_iso_datetime(parent_row.get("提交时间")),
                        ),
                    )
                    batch_id = int(cur.lastrowid)
                else:
                    batch_id = int(existed[0])
                    self.db.execute(
                        """
                        UPDATE reimbursement_batch
                        SET legacy_details_raw = ?,
                            legacy_project_name = ?,
                            legacy_claimant_name = ?,
                            updated_at = ?
                        WHERE id = ?
                        """,
                        (
                            details_text,
                            project_name,
                            claimant_name,
                            now_iso(),
                            batch_id,
                        ),
                    )

                for env_id, env_amount, _ in matched_envelopes:
                    self.db.execute(
                        """
                        INSERT OR IGNORE INTO reimbursement_batch_item
                          (batch_id, envelope_id, amount, status, created_at)
                        VALUES (?, ?, ?, 'active', ?)
                        """,
                        (batch_id, env_id, env_amount, now_iso()),
                    )

                item_stat = self.db.execute(
                    "SELECT COUNT(*) as cnt, IFNULL(SUM(amount), 0) as sum_amount FROM reimbursement_batch_item WHERE batch_id = ?",
                    (batch_id,),
                ).fetchone()
                item_count = int(item_stat[0]) if item_stat else 0
                item_sum = float(item_stat[1] or 0) if item_stat else 0

                self.db.execute(
                    """
                    UPDATE reimbursement_batch
                    SET item_count = ?,
                        total_amount = CASE WHEN IFNULL(total_amount, 0) <= 0 THEN ? ELSE total_amount END,
                        updated_at = ?
                    WHERE id = ?
                    """,
                    (item_count, item_sum, now_iso(), batch_id),
                )

            stats.success += 1 + len(child_rows)

        for row in rows:
            stats.total += 1
            legacy_data_id = as_text(row.get("data_id"))
            detail = as_text(row.get("报销详情"))

            if legacy_data_id:
                if current_parent is not None:
                    flush_parent(current_parent)
                current_parent = {"row": row, "child_rows": [], "details": []}
                if detail and detail != "流水号":
                    current_parent["details"].append(detail)
                continue

            if detail == "流水号":
                stats.success += 1
                continue

            if current_parent is None:
                stats.failed += 1
                self.error(
                    table,
                    row.get("__row_no"),
                    None,
                    "BATCH_ORPHAN_DETAIL",
                    "detail row appears before any parent row",
                    row,
                )
                continue

            if detail:
                current_parent["details"].append(detail)
            current_parent["child_rows"].append(row)

        if current_parent is not None:
            flush_parent(current_parent)

        self.finish_job(table, stats)


def load_mapping(path: Path) -> dict[str, Any]:
    if not path.exists():
        return {
            "default_userid": "legacy_import",
            "project_map": {},
            "user_map": {},
            "entity_map": {},
        }
    return json.loads(path.read_text(encoding="utf-8"))


def discover_local_db_path(repo_root: Path) -> Path | None:
    base = repo_root / "packages" / "api" / ".wrangler" / "state" / "v3" / "d1" / "miniflare-D1DatabaseObject"
    if not base.exists():
        return None
    candidates = sorted(base.glob("*.sqlite"), key=lambda p: p.stat().st_mtime, reverse=True)
    return candidates[0] if candidates else None


def main() -> None:
    parser = argparse.ArgumentParser(description="Legacy finance import for fin_data")
    parser.add_argument("--repo-root", default=str(Path(__file__).resolve().parents[3]))
    parser.add_argument("--input-dir", default="doc/legacy_data/fin_data")
    parser.add_argument("--mapping-file", default="doc/legacy_data/fin_data_mappings.template.json")
    parser.add_argument("--summary-file", default="doc/legacy_data/legacy_fin_import_summary.json")
    parser.add_argument("--mode", choices=["dry-run", "apply"], default="dry-run")
    parser.add_argument("--db-path", default="")
    parser.add_argument(
        "--tables",
        default="payee,entity,contract,payment_request,reimbursement_envelope,reimbursement_batch",
        help="comma-separated table groups to import: payee,entity,contract,payment_request,reimbursement_envelope,reimbursement_batch",
    )
    args = parser.parse_args()

    repo_root = Path(args.repo_root).resolve()
    input_dir = (repo_root / args.input_dir).resolve()
    mapping_file = (repo_root / args.mapping_file).resolve()
    summary_file = (repo_root / args.summary_file).resolve()
    mapping = load_mapping(mapping_file)

    db: sqlite3.Connection | None = None
    if args.mode == "apply":
        db_path = Path(args.db_path).resolve() if args.db_path else discover_local_db_path(repo_root)
        if db_path is None or not db_path.exists():
            raise RuntimeError("Local D1 sqlite file not found. Use --db-path to specify it.")
        db = sqlite3.connect(db_path)
        db.row_factory = sqlite3.Row

    migrator = Migrator(db=db, mapping=mapping, mode=args.mode)

    selected = {item.strip() for item in args.tables.split(",") if item.strip()}

    files = {
        "payee_external": input_dir / "常用收款人（外部）_20260228101032.xlsx",
        "payee_internal": input_dir / "常用收款人（内部）_20260228101145.xlsx",
        "entity": input_dir / "公司目录（我方）_20260228101307.xlsx",
        "contract": input_dir / "对外付款合同结算管理_20260228101109.xlsx",
        "payment_request": input_dir / "对外付款申请_20260228101021.xlsx",
        "reimbursement_envelope": input_dir / "个人报销预申请_20260228101235.xlsx",
        "reimbursement_batch": input_dir / "个人报销集中申报_20260228101252.xlsx",
    }

    for key, path in files.items():
        if not path.exists():
            raise FileNotFoundError(f"missing source file: {path}")

    if "payee" in selected:
        migrator.process_payee(load_rows(files["payee_external"]), "external", files["payee_external"].name)
        migrator.process_payee(load_rows(files["payee_internal"]), "internal", files["payee_internal"].name)
    if "entity" in selected:
        migrator.process_entity(load_rows(files["entity"]), files["entity"].name)
    if "contract" in selected:
        migrator.process_contract(load_rows(files["contract"]), files["contract"].name)
    if "payment_request" in selected:
        migrator.process_payment_request(load_rows(files["payment_request"]), files["payment_request"].name)
    if "reimbursement_envelope" in selected:
        migrator.process_reimbursement_envelope(load_rows(files["reimbursement_envelope"]), files["reimbursement_envelope"].name)
    if "reimbursement_batch" in selected:
        migrator.process_reimbursement_batch(load_rows(files["reimbursement_batch"]), files["reimbursement_batch"].name)

    summary_file.parent.mkdir(parents=True, exist_ok=True)
    summary_file.write_text(json.dumps(migrator.summary, ensure_ascii=False, indent=2), encoding="utf-8")

    if migrator.dropped_error_samples > 0:
        print(json.dumps({"warning": "error samples truncated", "dropped": migrator.dropped_error_samples}, ensure_ascii=False))

    if db is not None:
        db.commit()
        db.close()

    print(json.dumps({"ok": True, "mode": args.mode, "summary_file": str(summary_file)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
